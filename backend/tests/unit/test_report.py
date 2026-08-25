"""Report Agent(巡检记录渲染)的单元测试 —— 全程不联网。

重点不是覆盖率,是锁住三件「坏了就出错误文档」的事:
  · 文档里的隐患项/级别必须与 analyze 返回的契约数据**逐字一致**(留档材料的底线);
  · analyze 失败时**透传失败信封**,绝不渲染一份看起来正常的空文档;
  · 模型**没调渲染工具**时,不许让它编的「记录已出好,编号 …」流到用户面前
    (2026-08-10 真机验收抓获,防线是 RequireReceiptSource 结构件,见本文件最后一节)。
"""

from __future__ import annotations

import io
import re
from typing import Any

import pytest
from docx import Document
from langchain_core.messages import AIMessage, BaseMessage, ToolMessage
from langchain_core.outputs import ChatGeneration, ChatResult
from pydantic import Field

from gyt.agents.report import (
    REPORT_AGENT_NAME,
    REPORT_GIVE_UP_MESSAGE,
    REPORT_RETRY_NUDGE,
    build_report_agent,
)
from gyt.agents.report import tools as report_tools
from gyt.core import artifacts, llm
from gyt.core.errors import ErrorCode
from tests.unit.test_base_agent import FakeChatModel


class _FakeAnalyze:
    """替身:safety 的 ``_recognize`` —— **普通协程函数**,直接 await 调用。

    W9 之前这里替的是 ``analyze_site_photo`` 那个 BaseTool(``.ainvoke(payload)``)。
    换掉是因为那个工具现在还会把违规项**登记进隐患台账**,而 report 的保真复调是工具
    内部直调、拿不到 config —— 登记会落进 ``project_id=''``,同一张照片的同一个隐患在
    台账里变成两行、一行还无主(理由写在 report/tools.py 顶部「数据保真」那段)。
    ``seen`` 仍旧记成 ``{"artifact_id": …}`` 的形状,下面那些断言一个字都不用改。
    """

    def __init__(self, envelope: Any) -> None:
        self.envelope = envelope
        self.seen: list[dict[str, Any]] = []

    async def __call__(self, artifact_id: str) -> Any:
        self.seen.append({"artifact_id": artifact_id})
        return self.envelope


def _ok_envelope(**data: Any) -> dict[str, Any]:
    base = {
        "label": "violation",
        "violations": ["高空作业未系安全带", "未穿反光衣"],
        "severity": {"高空作业未系安全带": "重大", "未穿反光衣": "一般"},
        "max_severity": "重大",
        "note": "左侧工人腰部无挂钩;右侧远处一人看不清",
    }
    base.update(data)
    return {"ok": True, "data": base, "user_msg": "", "error_code": None}


def _patch(monkeypatch: pytest.MonkeyPatch, envelope: Any) -> _FakeAnalyze:
    fake = _FakeAnalyze(envelope)
    monkeypatch.setattr(report_tools, "_recognize", fake)
    return fake


async def _render(artifact_id: str = "a" * 32) -> dict[str, Any]:
    result = await report_tools.render_inspection_report.ainvoke({"artifact_id": artifact_id})
    assert isinstance(result, dict)
    return result


def _docx_text(report_id: str) -> str:
    """把落盘的 docx 读回来拼成一段文本,断言用。"""
    payload = artifacts.resolve(report_id).read_bytes()
    doc = Document(io.BytesIO(payload))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


async def test_文档内容与契约数据逐字一致(monkeypatch: pytest.MonkeyPatch) -> None:
    """留档材料的底线:隐患项、级别、备注一个字都不许走样。"""
    fake = _patch(monkeypatch, _ok_envelope())
    result = await _render("f" * 32)

    assert result["ok"] is True
    assert fake.seen == [{"artifact_id": "f" * 32}], "必须原样把编号交给 analyze"

    data = result["data"]
    text = _docx_text(data["report_id"])
    assert "工地安全巡检记录" in text
    assert "f" * 32 in text, "照片编号要进文档,追责时靠它对回原图"
    assert "高空作业未系安全带" in text and "重大" in text
    assert "未穿反光衣" in text and "一般" in text
    assert "左侧工人腰部无挂钩" in text, "note 是复核线索,必须落进文档"
    assert "持证安全员" in text, "初筛定位的免责句必须出现在每份产物里"
    assert data["max_severity"] == "重大"
    assert data["report_no"].startswith("GYT-")


async def test_登记为REPORT产物且可解析(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch(monkeypatch, _ok_envelope())
    result = await _render()
    meta = artifacts.read_meta(result["data"]["report_id"])
    assert meta["kind"] == "REPORT"
    assert meta["ext"] == ".docx"


async def test_合规照片渲染为未见隐患(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch(
        monkeypatch,
        _ok_envelope(label="compliant", violations=[], severity={}, max_severity=None, note=""),
    )
    result = await _render()
    text = _docx_text(result["data"]["report_id"])
    assert "未见明显隐患" in text
    assert "未发现受控清单内的隐患" in text
    assert "未发现受控清单内的隐患" in result["user_msg"]


async def test_非工地照片如实写进文档并提醒(monkeypatch: pytest.MonkeyPatch) -> None:
    """not_site 也要出记录(留下"这张不是现场"的痕迹),但话必须说明白。"""
    _patch(
        monkeypatch,
        _ok_envelope(
            label="not_site", violations=[], severity={}, max_severity=None, note="办公室场景"
        ),
    )
    result = await _render()
    text = _docx_text(result["data"]["report_id"])
    assert "非作业现场照片" in text
    assert "核对" in result["user_msg"]


async def test_识别失败时透传失败信封而不是渲染空文档(monkeypatch: pytest.MonkeyPatch) -> None:
    """**这条是底线。** 渲染一份"看起来正常"的空文档比不渲染危险得多 ——
    它会变成一份内容为「无隐患」的正式留档。"""
    failure = {
        "ok": False,
        "data": None,
        "user_msg": "没找到这张照片,请重新传一次。",
        "error_code": "NOT_FOUND",
    }
    _patch(monkeypatch, failure)
    result = await _render()
    assert result["ok"] is False
    assert result["error_code"] == ErrorCode.NOT_FOUND.value
    assert result["user_msg"] == "没找到这张照片,请重新传一次。"


async def test_词表外级别缺失时落为待定级(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch(
        monkeypatch, _ok_envelope(violations=["未佩戴安全帽"], severity={}, max_severity="待定级")
    )
    result = await _render()
    text = _docx_text(result["data"]["report_id"])
    assert "未佩戴安全帽" in text and "待定级" in text


def test_能组装出report_agent() -> None:
    agent = build_report_agent()
    assert agent.name == REPORT_AGENT_NAME


def test_提示词存在且含红线() -> None:
    from gyt.agents.report import REPORT_DIR
    from gyt.core.base_agent import load_prompt

    body = load_prompt(REPORT_DIR)
    assert "render_inspection_report" in body
    assert "原样照抄" in body
    assert "日报" in body, "必须明说日报还没有,别让它口头许诺"


# ---------------------------------------------------------------------------
# 英雄链(inspection)的结构 —— 那条硬边必须真的是图边,不是提示词里的许愿
# ---------------------------------------------------------------------------


def test_英雄链的safety到report是写死的图边() -> None:
    """「拍照识违规 → 自动出记录」交给 LLM 决策第二跳,演示当场会随机漏掉
    出报告那一环(方案 D 决策原文)。这条测试锁住:它是图边,不是概率。"""
    from gyt.graph import build_inspection_chain

    chain = build_inspection_chain()
    assert chain.name == "inspection"

    drawable = chain.get_graph()
    assert {"safety", "report"} <= set(drawable.nodes)

    edges = {(e.source, e.target) for e in drawable.edges}
    assert ("safety", "report") in edges, "硬边丢了 —— 出报告变回了 LLM 的心情"
    assert ("__start__", "safety") in edges
    assert ("report", "__end__") in edges


def test_inspection已入登记表且路由词表同步() -> None:
    """加 Agent 的规矩:AGENT_REGISTRY 与 scorers.ROUTING_AGENTS 一起改。
    只改一边的话,routing.csv 里写 inspection 会当场炸,或者路由永远派不到它。"""
    from eval.scorers import ROUTING_AGENTS

    from gyt.graph import AGENT_REGISTRY, INSPECTION_AGENT_NAME

    spec = next((s for s in AGENT_REGISTRY if s.name == INSPECTION_AGENT_NAME), None)
    assert spec is not None, "inspection 没挂进登记表,Supervisor 永远派不到英雄链"
    assert "巡检" in spec.summary and "记录" in spec.summary
    assert INSPECTION_AGENT_NAME in ROUTING_AGENTS


# ===========================================================================
# 「编号溯源」守卫(RequireReceiptSource)—— 硬边保证 report 被叫起来,
# 保证不了它真干活
#
# 2026-08-10 真机验收抓获:report 没调 render_inspection_report,直接照着
# prompt.md 里的示范句编了一段「巡检记录出好了,编号 GYT-…」,而磁盘上
# 没有那份 docx。上面那条 test_能组装出report_agent 只验了「能组装」——
# 守卫挂没挂上、挂上了管不管用,原来一条断言都没有,这一节补的就是这个洞。
#
# 写法上刻意**不去反查编译后的图里有没有那个中间件**(脆、且换个 langchain
# 小版本就得重写)。这里用一个「张口就报假编号」的假模型跑一遍**真实的**
# report agent,断言用户最终收到的是什么 —— 那才是这次改动要保证的东西。
# ===========================================================================

RENDER_TOOL_NAME = "render_inspection_report"
"""渲染工具的注册名。打回话术必须点到它,否则模型照着一个不存在的名字重试。"""

PHOTO_ID = "d" * 32
"""这一节用的照片编号。工具内部的 analyze 已被替身接管,不需要真有这张照片。"""

FABRICATED_REPORT_LINE = "巡检记录出好了,编号 GYT-20260809-153739。"
"""真机那次编出来的原话。它一个字都不许出现在最终消息里 —— 这个编号在磁盘上
不存在,工友拿着它去取件只会扑空。"""

HONEST_REPORT_LINE = "巡检记录出好了,这次共 2 处隐患,最高级别是重大。"
"""对照组:模型**真调过工具之后**说的话。守卫不许把这种话也一起吃掉。"""


class _ScriptedChatModel(FakeChatModel):
    """按剧本答话的假模型:记下每一次被调用时看到的最后一条消息。

    继承 test_base_agent 那份 FakeChatModel 是为了白拿它的 bind_tools ——
    langgraph-prebuilt 会去读绑定结果里每个工具的 .get("name"),直接塞 BaseTool
    对象会炸,这个坑那边已经踩平并写了注释,不必在这里抄第二遍(DRY)。

    tool_artifact_id 为 None = 「首答永远不调工具」的病模型;
    给了编号 = 首答就调渲染工具的正常模型(拿到工具结果后再出终答)。
    """

    calls: list[str] = Field(default_factory=list)
    tool_artifact_id: str | None = None

    def _generate(
        self,
        messages: list[BaseMessage],
        stop: list[str] | None = None,
        run_manager: Any = None,
        **kwargs: Any,
    ) -> ChatResult:
        # 记最后一条:重试那次的末尾正是守卫追加的系统校验,断言要靠它。
        self.calls.append(str(messages[-1].content))
        # 按「上文里有没有工具结果」分轮,而不是按调用次数 —— 打回重试也会让
        # 计数往前走,按次数分轮的话第二轮会错判成终答轮。
        already_used_tool = any(isinstance(m, ToolMessage) for m in messages)
        if self.tool_artifact_id and not already_used_tool:
            message = AIMessage(
                content="",
                tool_calls=[
                    {
                        "name": RENDER_TOOL_NAME,
                        "args": {"artifact_id": self.tool_artifact_id},
                        "id": "call-report-1",
                    }
                ],
            )
        else:
            message = AIMessage(content=self.reply_text)
        return ChatResult(generations=[ChatGeneration(message=message)])


def _stub_model(monkeypatch: pytest.MonkeyPatch, model: _ScriptedChatModel) -> None:
    """把假模型塞给真实工厂。

    打在 llm 模块对象上:base_agent 用的是 `from gyt.core import llm` +
    `llm.get_chat_model(...)`,属性查找发生在调用时,在这里 setattr 才生效。
    """

    def _fake_get_chat_model(purpose: str = "text", **_overrides: Any) -> _ScriptedChatModel:
        return model

    monkeypatch.setattr(llm, "get_chat_model", _fake_get_chat_model)


async def _run_report(ask: str) -> list[BaseMessage]:
    """跑一遍真实的 report agent(真提示词、真工具、真中间件),返回最终消息列表。

    刻意不收 model 形参:假模型是靠 _stub_model 打在 llm 模块上的,
    收一个用不到的形参会让后人以为「传进去就生效」,拿真 ChatOpenAI 去跑。
    """
    agent = build_report_agent()
    result = await agent.ainvoke({"messages": [("user", ask)]})
    return list(result["messages"])


async def test_模型不调工具时用户收到的是没出成而不是假编号(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """**这条是这次改动的验收断言。**

    「没出成」工友看得见、也重试得了;「以为出成了」他揣着一个不存在的编号
    走人,后面没有任何一环能救回来 —— 所以 report 的 on_give_up 是 fail。
    """
    model = _ScriptedChatModel(reply_text=FABRICATED_REPORT_LINE)
    _stub_model(monkeypatch, model)

    messages = await _run_report(f"给照片 {PHOTO_ID} 出份巡检记录")

    texts = [str(m.content) for m in messages]
    assert texts[-1] == REPORT_GIVE_UP_MESSAGE, "没落盘就得说没落盘,不许放模型编的话出去"
    assert all(FABRICATED_REPORT_LINE not in text for text in texts), (
        "编出来的假编号连一个字都不许留在消息里 —— 前端会照样把它渲染给工友看"
    )
    assert len(model.calls) == 2, "只重试一次;守卫绝不能自己造第二个循环(熔断归上层)"
    assert "系统校验" in model.calls[1], "重试那一次必须真把打回话术喂给模型"


async def test_首答就调工具时守卫不插手(monkeypatch: pytest.MonkeyPatch) -> None:
    """正常路径的护栏:守卫只拦「没调工具」,不许顺手把真干过活的回答也改掉。

    这条同时证明拦截发生在**回合首答**这一环:拿到工具结果之后的那次纯文本
    终答不是首答,不该被打回(打回它反而制造循环)。
    """
    fake_analyze = _patch(monkeypatch, _ok_envelope())
    model = _ScriptedChatModel(reply_text=HONEST_REPORT_LINE, tool_artifact_id=PHOTO_ID)
    _stub_model(monkeypatch, model)

    messages = await _run_report(f"给照片 {PHOTO_ID} 出份巡检记录")

    assert fake_analyze.seen == [{"artifact_id": PHOTO_ID}], "工具得真跑起来,不然测的不是这条路"
    texts = [str(m.content) for m in messages]
    assert texts[-1] == HONEST_REPORT_LINE, "真调过工具的回答必须原样送到用户面前"
    assert REPORT_GIVE_UP_MESSAGE not in texts, "活干成了却报「没出成」,比不拦还糟"
    assert len(model.calls) == 2, "一次要工具、一次出终答,中间不该有打回"


def test_给工友的失败话术不许回退成技术腔() -> None:
    """这句话是本次改动里**唯一**会出现在工地师傅眼前的字符串,单独钉住。

    反面教材(以前的写法就是这一类):「渲染失败,请重试」「工具调用异常」——
    师傅看不懂"渲染"是什么,也不知道该重试什么、找谁。
    """
    assert re.search(r"[A-Za-z]", REPORT_GIVE_UP_MESSAGE) is None, (
        "一个英文字母都不该有:类名、工具名、错误码全是内部词"
    )
    for forbidden in ("编号", "已生成", "生成成功", "记录已出"):
        assert forbidden not in REPORT_GIVE_UP_MESSAGE, (
            f"「{forbidden}」会让人以为文档已经出来了 —— 这正是这次要防的事"
        )
    assert "没出成" in REPORT_GIVE_UP_MESSAGE, "得先把「这次没成」说明白"
    assert "再发一遍" in REPORT_GIVE_UP_MESSAGE, "得给一个当场能做的动作,不能只报丧"


def test_打回话术点名的工具真的存在() -> None:
    """措辞随便改,但点名的那个工具名不能变成不存在的东西 ——
    模型照着一个查无此名的工具重试,这次打回就白费了。"""
    tool_names = {tool.name for tool in report_tools.REPORT_TOOLS}
    assert RENDER_TOOL_NAME in tool_names, "渲染工具改名了,下面这条断言的前提就没了"
    assert RENDER_TOOL_NAME in REPORT_RETRY_NUDGE, "打回话术必须点名该调哪个工具"


async def test_英雄链上safety之后的首答也必须受管(monkeypatch: pytest.MonkeyPatch) -> None:
    """复刻 2026-08-10 真机那一幕的**真实消息形状**,而不是实验室里的形状。

    上面那条 test_模型不调工具时用户收到的是没出成而不是假编号 是用一条
    HumanMessage 直接叫醒 report 的 —— 那条路在生产上不存在。生产上 report
    永远是被硬边从 safety 推过来的,它首答之前看到的最后一条是这样:

        HumanMessage(用户)
        AIMessage(safety, tool_calls=[analyze_site_photo])
        ToolMessage(analyze_site_photo)
        AIMessage(safety, "发现 1 处重大隐患……")   ← 首答之前的最后一条

    这个形状是实测出来的(把假模型塞进复刻的 safety→report 子图,打印
    report 模型每次被调时看到的消息序列),不是照着代码推的。
    """
    model = _ScriptedChatModel(reply_text=FABRICATED_REPORT_LINE)
    _stub_model(monkeypatch, model)
    agent = build_report_agent()

    result = await agent.ainvoke(
        {
            "messages": [
                ("user", f"看看照片 {PHOTO_ID},顺便出份巡检记录"),
                AIMessage(
                    content="",
                    name="safety",
                    tool_calls=[
                        {
                            "name": "analyze_site_photo",
                            "args": {"artifact_id": PHOTO_ID},
                            "id": "call-safety-1",
                        }
                    ],
                ),
                ToolMessage(
                    content='{"ok": true}',
                    tool_call_id="call-safety-1",
                    name="analyze_site_photo",
                ),
                AIMessage(content="这张照片发现 1 处重大隐患:高空作业未系安全带。", name="safety"),
            ]
        }
    )

    texts = [str(m.content) for m in result["messages"]]
    assert texts[-1] == REPORT_GIVE_UP_MESSAGE, "英雄链才是 report 唯一的来路,这条路必须也拦得住"
    assert all(FABRICATED_REPORT_LINE not in text for text in texts)


# ---------------------------------------------------------------------------
# 零误伤护栏:prompt.md 明文要求的两种「该纯文本回答」的回合
#
# 2026-08-10 复核抓到的第二条 HIGH:初版守卫的判据是「回合首答必须调工具」,
# 那会把下面这两种**正确行为**也一起顶替掉 —— 反问被吃掉就是对话死锁。
# 改成编号溯源之后它们天然不受影响(压根不含编号),这两条钉住这个性质,
# 防的是将来有人图省事把判据换回「首答必须调工具」。
# ---------------------------------------------------------------------------

LEGIT_ASK_WHICH = "师傅,给哪张出记录?"
"""prompt.md 第 23 行明文要求的反问。"""

LEGIT_REFUSE_DAILY = "日报功能还没上线,现在能出单张照片的巡检记录。"
"""prompt.md 第 45 行明文要求的如实拒答。"""


@pytest.mark.parametrize(
    ("reply", "label"),
    [(LEGIT_ASK_WHICH, "拿不准是哪张照片时的反问"), (LEGIT_REFUSE_DAILY, "问日报时的如实拒答")],
)
async def test_合法的纯文本回合不许被守卫顶替(
    monkeypatch: pytest.MonkeyPatch, reply: str, label: str
) -> None:
    """这两种回答不调工具、也不报编号 —— 守卫必须一个字都不动。"""
    model = _ScriptedChatModel(reply_text=reply)
    _stub_model(monkeypatch, model)

    messages = await _run_report(f"给照片 {PHOTO_ID} 出份巡检记录")

    texts = [str(m.content) for m in messages]
    assert texts[-1] == reply, f"{label}被顶替了 —— 这是 prompt.md 要求的正确行为"
    assert REPORT_GIVE_UP_MESSAGE not in texts, "没有编造编号,不该走到失败话术"


async def test_念真回执时守卫放行(monkeypatch: pytest.MonkeyPatch) -> None:
    """真调过工具、编号有出处 —— 守卫不许把这种话也吃掉。

    这条和上面「编造被顶替」是一对:只证明「会拦」不够,还得证明「不乱拦」。
    """
    model = _ScriptedChatModel(reply_text=HONEST_REPORT_LINE, tool_artifact_id=PHOTO_ID)
    _stub_model(monkeypatch, model)

    messages = await _run_report(f"给照片 {PHOTO_ID} 出份巡检记录")

    texts = [str(m.content) for m in messages]
    assert REPORT_GIVE_UP_MESSAGE not in texts, "真调过工具还被顶替,守卫太凶了"


# =============================================================================
# 自定义标题(2026-08-24 加,复验报告 P0-2)
#
# 这一组守的不是「标题能不能改」,而是**改了之后文种还在不在**。
# 完整推演在 agents/report/tools.py 的 _clean_title 头注,一句话:
# 标题属于用户(这一轮巡检叫什么),文种属于系统(它是一份什么文件),
# 一份文档不许因为标题就变成另一个文种。
# =============================================================================


async def _render_titled(title: str, artifact_id: str = "b" * 32) -> dict[str, Any]:
    """带自定义标题走一遍工具。**经 ainvoke 而不是直调内部函数** —— 要连
    工具签名(模型看得见的那一层)一起验,签名漏了参数这组测试才该红。"""
    result = await report_tools.render_inspection_report.ainvoke(
        {"artifact_id": artifact_id, "title": title}
    )
    assert isinstance(result, dict)
    return result


def _docx_title(report_id: str) -> str:
    """只取文档标题(docgen 用 add_heading 放的那个段落),不含表格。

    不能复用 `_docx_text`:它把段落和表格拼在一起,而这一组要分清
    「标题」与「元信息表里的文档类型」——它俩正是本组要证明不相等的两样东西。
    """
    payload = artifacts.resolve(report_id).read_bytes()
    doc = Document(io.BytesIO(payload))
    for para in doc.paragraphs:
        if para.text.strip():
            return para.text.strip()
    raise AssertionError("文档里一个非空段落都没有")


async def test_不传标题时用默认标题(monkeypatch: pytest.MonkeyPatch) -> None:
    """没人要求标题就别自作主张 —— 保持 2026-08-24 之前的行为逐字不变。"""
    _patch(monkeypatch, _ok_envelope())
    result = await _render("c" * 32)
    assert _docx_title(result["data"]["report_id"]) == report_tools.DEFAULT_REPORT_TITLE


async def test_传了标题就用它(monkeypatch: pytest.MonkeyPatch) -> None:
    """复验报告 P0-2 要的就是这一条:用户说的标题真的进了文件,不是只在聊天里回显。"""
    _patch(monkeypatch, _ok_envelope())
    result = await _render_titled("海之子验收测试")
    assert _docx_title(result["data"]["report_id"]) == "海之子验收测试"


async def test_自定义标题改不掉文种(monkeypatch: pytest.MonkeyPatch) -> None:
    """🔴 本组最值钱的一条。

    标题被写成另一个文种的名字时,元信息表里那行「文档类型」必须还是
    「工地安全巡检记录」—— 否则一份 AI 初筛记录可以被叫成
    「现场无隐患确认书」拿去交差,而文档本身不提供任何反证。
    """
    _patch(monkeypatch, _ok_envelope())
    result = await _render_titled("现场无隐患确认书")
    report_id = result["data"]["report_id"]

    assert _docx_title(report_id) == "现场无隐患确认书", "标题该听用户的"
    text = _docx_text(report_id)
    assert "文档类型" in text and report_tools.DEFAULT_REPORT_TITLE in text, (
        "文种必须留在元信息表里,这是标题可自定义的前提"
    )


async def test_标题里的折行与控制字符被压平(monkeypatch: pytest.MonkeyPatch) -> None:
    """docx 的标题是单行段落:塞进 \\n 不报错,只渲染成一个看不见的怪空格,
    而人对着 Word 找不出哪里不对。所以要在进文档之前压平。"""
    _patch(monkeypatch, _ok_envelope())
    result = await _render_titled("海之子\n验收\t测试\x07报告")
    assert _docx_title(result["data"]["report_id"]) == "海之子 验收 测试 报告"


async def test_连续空白压成一个(monkeypatch: pytest.MonkeyPatch) -> None:
    """粘贴产物常带一串空格,渲染出来像排版事故。"""
    _patch(monkeypatch, _ok_envelope())
    result = await _render_titled("  海之子   验收测试  ")
    assert _docx_title(result["data"]["report_id"]) == "海之子 验收测试"


async def test_纯空白标题退回默认(monkeypatch: pytest.MonkeyPatch) -> None:
    """「传了个纯空格」和「没传」在意图上没区别 —— 不为标题让整份文档生不出来。"""
    _patch(monkeypatch, _ok_envelope())
    result = await _render_titled("   \n\t  ")
    assert _docx_title(result["data"]["report_id"]) == report_tools.DEFAULT_REPORT_TITLE


async def test_超长标题被截断(monkeypatch: pytest.MonkeyPatch) -> None:
    """排版边界,不是安全边界:超了会折行、把元信息表挤到第二页。"""
    _patch(monkeypatch, _ok_envelope())
    long_title = "验" * 100
    result = await _render_titled(long_title)
    title = _docx_title(result["data"]["report_id"])
    assert len(title) == report_tools._TITLE_MAX_LEN
    assert title == "验" * report_tools._TITLE_MAX_LEN


def test_净化函数本身(monkeypatch: pytest.MonkeyPatch) -> None:
    """纯函数直测:上面那几条走完整渲染,这条钉边界值,坏了能一眼看出是哪一步。"""
    clean = report_tools._clean_title
    assert clean("") == report_tools.DEFAULT_REPORT_TITLE
    assert clean("正常标题") == "正常标题"
    assert clean("a b") == "a b", "不间断空格也算空白"
    assert len(clean("字" * 200)) == report_tools._TITLE_MAX_LEN
