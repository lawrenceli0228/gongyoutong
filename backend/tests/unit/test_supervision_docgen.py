"""共用 docx 骨架的单元测试(W9 S2)—— 全程不联网、不落盘、不碰库。

这一份钉的是「渲染出来的文书坏了但看不出来」那一类事:

  · **顺序与结构**:标题 → 元信息表 → 各节 → 签字栏 → 免责句。少一节、错一位,
    打开文件的人不会觉得是 bug,只会觉得这份文书本来就长这样。
  · **表格列数对不上要当场炸**,不许把内容错位到相邻列 —— 错位的隐患级别
    (「重大」落进「隐患项」列)在留档文书上看不出问题,却是要追责的材料。
  · **两句免责句不许互相顶替**:监理文书那句写着「总监签字后生效」,巡检记录
    那句写着「以持证安全员现场判定为准」,套错等于对外宣称了一件没发生的事。
    结构上的保证是 ``disclaimer`` 必填无默认,这里用「不传就 TypeError」钉住。
  · **不落盘**:方案 §6.4 的原子性(先渲染到内存 → register → 一个事务写库)
    全靠渲染阶段不碰文件系统;自己落盘就会在②之前留下认不出来的半成品。
"""

from __future__ import annotations

import dataclasses
import io
import re
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml.ns import qn

from gyt.agents.supervision import docgen
from gyt.agents.supervision.docgen import (
    SUPERVISION_DISCLAIMER,
    SUPERVISION_SIGNATURE_LINES,
    TableSection,
    TextSection,
    render_document,
)
from gyt.core.doc_no import NOTICE_NO_PATTERN, DocKind, new_doc_no

# --- 读回工具 ---------------------------------------------------------------


def _flow(payload: bytes) -> list[tuple[str, Any]]:
    """把 docx 读回成「文档流」:按正文里的真实先后顺序列出段落与表格。

    为什么不直接用 ``doc.paragraphs`` + ``doc.tables``:那是两条各自独立的列表,
    只能证明「有这些东西」,证明不了「表格排在哪两段之间」。而这份骨架要保证的
    恰恰是顺序 —— 免责句跑到正文前面去了,靠前两个列表一条都测不出来。
    """
    doc = Document(io.BytesIO(payload))
    paragraphs = iter(doc.paragraphs)
    tables = iter(doc.tables)
    flow: list[tuple[str, Any]] = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            flow.append(("p", next(paragraphs).text))
        elif child.tag == qn("w:tbl"):
            table = next(tables)
            flow.append(("t", [[cell.text for cell in row.cells] for row in table.rows]))
    return flow


def _texts(payload: bytes) -> list[str]:
    return [value for kind, value in _flow(payload) if kind == "p"]


def _tables(payload: bytes) -> list[list[list[str]]]:
    return [value for kind, value in _flow(payload) if kind == "t"]


# --- 样例数据 ---------------------------------------------------------------

META = [("文书编号", "GYT-TZ-20260815-083000-abcd"), ("签发时间", "2026-08-15 08:30:00")]
HAZARD_TABLE = TableSection(
    heading="隐患明细",
    header=("序号", "隐患项", "级别"),
    rows=(("1", "高空作业未系安全带", "重大"),),
    empty_note="本次未发现受控清单内的隐患。",
)


def _render(**overrides: Any) -> bytes:
    kwargs: dict[str, Any] = {
        "title": "监理通知单",
        "meta": META,
        "sections": [TextSection(heading="整改要求", body="立即停止该部位作业。"), HAZARD_TABLE],
        "disclaimer": SUPERVISION_DISCLAIMER,
        "signature_lines": SUPERVISION_SIGNATURE_LINES,
    }
    kwargs.update(overrides)
    return render_document(**kwargs)


class Test骨架结构:
    def test_五段式顺序(self) -> None:
        """标题 → 元信息表 → 各节 → 签字栏 → 免责句,一个都不许错位。"""
        flow = _flow(_render())
        kinds = [kind for kind, _ in flow]
        assert kinds == ["p", "t", "p", "p", "p", "t", "p", "p", "p", "p"], kinds

        assert flow[0] == ("p", "监理通知单")
        assert flow[1][1] == [
            ["文书编号", "GYT-TZ-20260815-083000-abcd"],
            ["签发时间", "2026-08-15 08:30:00"],
        ]
        assert flow[2][1] == "整改要求"
        assert flow[3][1] == "立即停止该部位作业。"
        assert flow[4][1] == "隐患明细"
        assert flow[5][1] == [["序号", "隐患项", "级别"], ["1", "高空作业未系安全带", "重大"]]
        assert [value for _, value in flow[6:9]] == list(SUPERVISION_SIGNATURE_LINES)
        assert flow[9][1] == SUPERVISION_DISCLAIMER

    def test_元信息为空时不画空表(self) -> None:
        """一张零行的表在 Word 里是个诡异的细条,不如干脆不画。"""
        assert _tables(_render(meta=())) == [
            [["序号", "隐患项", "级别"], ["1", "高空作业未系安全带", "重大"]]
        ]

    def test_没有签字栏时正文与免责句直接相接(self) -> None:
        """巡检记录走的就是这条路 —— 它没有「签字后生效」这一环。"""
        texts = _texts(_render(signature_lines=()))
        assert texts[-1] == SUPERVISION_DISCLAIMER
        assert all(line not in texts for line in SUPERVISION_SIGNATURE_LINES)

    def test_免责句字号比正文小(self) -> None:
        doc = Document(io.BytesIO(_render()))
        tail = doc.paragraphs[-1]
        assert tail.text == SUPERVISION_DISCLAIMER
        assert tail.runs[0].font.size is not None and tail.runs[0].font.size.pt == 9


class Test节的两种形态:
    def test_文本节正文为空时只出小标题(self) -> None:
        texts = _texts(_render(sections=[TextSection(heading="处理意见")], signature_lines=()))
        assert "处理意见" in texts
        # 小标题之后紧接着就是免责句 —— 中间没有一段空段落。
        assert texts[texts.index("处理意见") + 1] == SUPERVISION_DISCLAIMER

    def test_表格节无数据时改画一句话而不是空表(self) -> None:
        """只有表头的空表读起来像「这一栏还没填」,而「本次未发现」是个明确结论。
        在追责场合这两件事意思完全不同。"""
        empty = dataclasses.replace(HAZARD_TABLE, rows=())
        payload = _render(meta=(), sections=[empty])
        assert _tables(payload) == [], "画了空表 —— 读的人会以为内容漏填了"
        assert "本次未发现受控清单内的隐患。" in _texts(payload)

    def test_表格节无数据且无兜底话时只剩小标题(self) -> None:
        bare = dataclasses.replace(HAZARD_TABLE, rows=(), empty_note="")
        payload = _render(meta=(), sections=[bare])
        assert _tables(payload) == []
        assert "隐患明细" in _texts(payload)

    def test_认不出的节类型当场炸而不是静默跳过(self) -> None:
        """悄悄少一节的文书,外观上完全正常 —— 这是最难发现的一种坏法。"""
        with pytest.raises(TypeError, match="认不出的节类型"):
            _render(sections=["随手传了个字符串"])


class Test错位防线:
    def test_行的格数与表头对不上就拒绝渲染(self) -> None:
        """**这条是错位防线。** 少一格会让「重大」挪进「隐患项」列,而打开文件
        的人看到的是一份格式完好、内容错位的正式文书。"""
        bad = dataclasses.replace(HAZARD_TABLE, rows=(("1", "高空作业未系安全带"),))
        with pytest.raises(ValueError, match="列数对不上") as caught:
            _render(sections=[bad])
        message = str(caught.value)
        assert "隐患明细" in message and "第 1 行" in message, "得说清是哪一节的哪一行"

    def test_报错点的是第一行出问题的那行(self) -> None:
        bad = dataclasses.replace(
            HAZARD_TABLE,
            rows=(("1", "甲", "重大"), ("2", "乙"), ("3", "丙")),
        )
        with pytest.raises(ValueError, match="第 2 行"):
            _render(sections=[bad])

    def test_列数对不上时一格都不许落进文档(self) -> None:
        """校验必须在建表之前做完,否则留下一张画到一半的残表 ——
        而调用方接住异常之后未必会丢掉这份 doc。"""
        bad = dataclasses.replace(HAZARD_TABLE, rows=(("1", "甲"),))
        with pytest.raises(ValueError):
            _render(sections=[bad])
        # 同一批数据改对之后照样能渲染,证明上一次没留下脏状态。
        assert _tables(_render()) != []

    def test_表格节没有表头就拒绝(self) -> None:
        bad = dataclasses.replace(HAZARD_TABLE, header=())
        with pytest.raises(ValueError, match="没有表头"):
            _render(sections=[bad])


class Test两句免责句不许混:
    def test_不传免责句直接是参数错误(self) -> None:
        """``disclaimer`` 必填且无默认值 —— 这是「漏传于是悄悄套上另一句」在
        结构上不可能发生的原因。有人哪天给它加个默认值,这条当场红。"""
        with pytest.raises(TypeError):
            render_document(title="监理通知单")  # type: ignore[call-arg]

    def test_免责句为空也拒绝(self) -> None:
        """空串等于对外宣称「这是终稿」。"""
        with pytest.raises(ValueError, match="免责句"):
            _render(disclaimer="   ")

    def test_监理文书那句与巡检记录那句不是同一句(self) -> None:
        """两句话讲的是两件事:一句是「签字后才生效」,一句是「AI 初筛、以持证
        安全员现场判定为准」。合并成一句的话,两边各自丢掉一半的边界声明。"""
        from gyt.agents.report.tools import _DISCLAIMER as REPORT_DISCLAIMER

        assert SUPERVISION_DISCLAIMER != REPORT_DISCLAIMER
        assert "持证安全员" in REPORT_DISCLAIMER and "持证安全员" not in SUPERVISION_DISCLAIMER
        assert "总监理工程师" in SUPERVISION_DISCLAIMER and "总监理工程师" not in REPORT_DISCLAIMER

    def test_监理免责句写清了D15的定位(self) -> None:
        """D15:文书是「AI 出稿,经总监理工程师签字后生效」,不是演示件。"""
        for word in ("工友通 AI", "总监理工程师", "签字", "正式文件"):
            assert word in SUPERVISION_DISCLAIMER, word

    def test_签字栏留空(self) -> None:
        """系统替人签字等于伪造 —— 签字这个动作本身就是「有人认下了这份文书」
        的唯一凭据,留空的横线由签字的人自己写。

        判据是「一个数字、一个字母都不许有」:自动填上的人名会带字母(拼音)或
        直接是汉字姓名,自动填上的日期一定带数字。年月日那三个字是**占位**,
        不是被填过的值,所以不算。
        """
        for line in SUPERVISION_SIGNATURE_LINES:
            assert re.search(r"[0-9A-Za-z]", line) is None, line
        assert SUPERVISION_SIGNATURE_LINES[0].endswith(":"), "签字行冒号后面得空着"
        assert SUPERVISION_SIGNATURE_LINES[1].endswith(":"), "盖章行冒号后面得空着"


class Test纯函数性质:
    def test_返回字节且不落盘(self, monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
        """§6.4 的原子性(渲染到内存 → register → 一个事务写库)全靠这条。
        自己落盘的话,第二份渲染炸掉时第一份已经躺在磁盘上,既没进注册表也没
        进库,连清理器都不认识它。"""
        monkeypatch.chdir(tmp_path)
        payload = _render()
        assert isinstance(payload, bytes) and payload[:2] == b"PK", "docx 是个 zip,头两字节应是 PK"
        assert list(tmp_path.iterdir()) == [], "渲染过程往磁盘上写东西了"

    def test_不改入参(self) -> None:
        meta = list(META)
        sections = [TextSection(heading="整改要求", body="立即停止该部位作业。"), HAZARD_TABLE]
        meta_before, sections_before = list(meta), list(sections)

        render_document(
            title="监理通知单", meta=meta, sections=sections, disclaimer=SUPERVISION_DISCLAIMER
        )

        assert meta == meta_before and sections == sections_before

    def test_节是冻结的(self) -> None:
        """文书内容是留档材料,渲染路上被就地改掉的话,写库写进去的是改过的版本,
        而两边对不上要到有人打开文件核对时才发现。"""
        with pytest.raises(dataclasses.FrozenInstanceError):
            HAZARD_TABLE.heading = "换个名字"  # type: ignore[misc]
        with pytest.raises(dataclasses.FrozenInstanceError):
            TextSection(heading="甲").body = "乙"  # type: ignore[misc]

    def test_同样的入参渲染两次内容一致(self) -> None:
        """骨架里不许偷偷塞时间戳一类的东西 —— 塞了就没法对比两份文书的差异。"""
        assert _flow(_render()) == _flow(_render())


class Test标题:
    def test_标题为空时拒绝(self) -> None:
        with pytest.raises(ValueError, match="标题"):
            _render(title="  ")


class Test与编号生成器合起来用:
    def test_一份完整的监理通知单(self) -> None:
        """冒烟:doc_no 摇号 → docgen 出稿。两个新件的接缝走一遍。"""
        no = new_doc_no(DocKind.NOTICE)
        payload = docgen.render_document(
            title="监理通知单",
            meta=(("文书编号", no), ("受检单位", "某某施工单位")),
            sections=(
                TextSection(heading="事由", body="现场巡检发现下列隐患,请限期整改。"),
                HAZARD_TABLE,
            ),
            disclaimer=SUPERVISION_DISCLAIMER,
            signature_lines=SUPERVISION_SIGNATURE_LINES,
        )
        flat = "\n".join(str(value) for _, value in _flow(payload))
        assert re.search(NOTICE_NO_PATTERN, flat), "文书编号得真的印在文档里"
        assert "高空作业未系安全带" in flat and "重大" in flat
