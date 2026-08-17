"""文书正文(``agents/supervision/documents.py``)的单元测试 —— 纯函数级,不起 HTTP、不碰库。

与另外两份的分工,别互相重复造:

    test_supervision_docgen.py  **骨架**:标题→元信息表→各节→签字栏→免责句的顺序、
                                表格列数对不上要炸、``disclaimer`` 必填、**不落盘**。
    test_supervision_api.py     **整条链**:七个端点、三条硬拦、三份文书原子产出、
                                落盘顺序、编号撞库重试。它也真渲染 docx,但只从
                                《工程暂停令》那一份上抽查了「编号印上了 + 免责句没串」。
    本文件                       **每种文书写了什么字**:五种各自的正文里那些
                                「不报错但内容错」的分支 —— 端到端那一层抽查不到,
                                因为它们要么要一个特定形状的证据链,要么根本不产生异常。

这一层钉的是这几类静默错误:

  · **免责句/签字栏串档或漏盖** —— 五种全测。只测暂停令那一份的话,新加的文书
    走岔了没人知道;而一份没写「未签字不算数」的停工文书看起来最正常。
  · **元信息表少印文书编号** —— 编号是这份纸的对外身份(电话里报它、排证据链按它),
    印不上去不会有任何异常,只是这份文书从此对不回台账。
  · **复工令引用不到原《工程暂停令》** —— ``_find_doc_no`` 找不到时给的是
    「(未找到)」这句话。真留空的话,复工令上会出现「原《工程暂停令》编号 。」
    这种断句,而它是要拿去存档的。
  · **空证据链画成一张空表** —— 空表在留档文件里读起来像「这一栏还没填」,
    而「暂无」是个明确结论;追责场合这两者意思完全不同。
  · **``doc_type`` 反查漏档** —— 证据链表里冒出一个英文 ``owner_report``。
  · **(F5)正文说的话附件兜不住** —— 《致建设单位报告》正文写死「已就此签发
    《监理通知单》与《工程暂停令》」,而它自己附的证据链一份都列不出来
    (``_context()`` 在 ``_sign()`` 之前建,本批三份还没入库)。正文说签了、
    附件说没有 —— 一份自相矛盾的法律文书,而这不会有任何报错。
    钉它的是 ``test_正文说已签发的文书必须在证据链里列得出来``,**必须带上
    「证据链为空」那一档**:非空那一档下,连出 bug 的那句原话都是绿的。
  · **(F4)复查照片取不出来** —— 照片编号存在 ``hazard_docs.photo_id`` 里,
    而文书上一直印的是 ``hazards.photo_id``(**首次发现**那张)。于是「复查必须
    挂一张照片」退化成一道提交时的门槛:提交完谁也取不出来,事后追责举不出证。
    连带的坑是两个 photo_id 在纸上叫同一个名字 —— 那等于拿发现时的照片
    当"整改后"的证据。

⚠️ 这些正文是**留档材料**。改这里的期望值之前先想清楚已经发出去的那些纸怎么办 ——
2026-08-16 从 supervision_api.py 切出来那次是纯搬家,判据就是「渲染出来的 docx
逐字节相同」,不是「测试绿」。
"""

from __future__ import annotations

import io
import re

import pytest
from docx import Document

from gyt.agents.supervision.docgen import SUPERVISION_DISCLAIMER, SUPERVISION_SIGNATURE_LINES
from gyt.agents.supervision.documents import (
    EVIDENCE_HEADING,
    FOUND_PHOTO_LABEL,
    ISSUING_KINDS,
    OWNER_REPORT_BATCH,
    REINSPECT_PHOTO_LABEL,
    UNASSIGNED_PROJECT_ZH,
    DocContext,
    render_doc,
)
from gyt.core.doc_no import DOC_TITLE_ZH, DocKind
from gyt.db import hazards

_KIND_IDS = [kind.name for kind in ISSUING_KINDS]
"""parametrize 的用例名。用 ``DocKind`` 的成员名,报错时一眼看出是哪种文书。"""

_HAZARD_NO = "GYT-H-20260816-093000-ab12"
_PHOTO_ID = "f" * 32
"""隐患**首次发现**那张照片(``hazards.photo_id``)。下面那个是复查那张,别混用。"""

_REINSPECT_PHOTO_ID = "c" * 32
"""某一次复查那张照片(``hazard_docs.photo_id``)。与 ``_PHOTO_ID`` 刻意取不同的值 ——
两者相等的话,「文书上印的到底是哪一张」这件事就测不出来了。"""

_SUSPENSION_NO = "GYT-ZT-20260816-093000-2222"

_带证据链的三种 = (DocKind.OWNER_REPORT, DocKind.RESUMPTION, DocKind.AUTHORITY_REPORT)
"""正文里带「处置经过」表的三种。另外两种(通知单、暂停令)是签发链的**第一步**,
那会儿名下本来就还没有任何文书 —— 给它们摆一张必然为空的表纯属噪音。"""


# --- 造素材(全是 NamedTuple,构造它们一次库都不用打)-------------------------


def _row(**overrides: object) -> hazards.HazardRow:
    """一条已定级、已停工的隐患。字段全给齐,免得哪天加列了这里静默走默认值。"""
    base: dict[str, object] = dict(
        id=7,
        hazard_no=_HAZARD_NO,
        project_id="gyt-a3",
        photo_sha256="0" * 64,
        photo_id=_PHOTO_ID,
        item="临边无防护",
        severity="重大",
        grade=hazards.GRADE_SEVERE,
        grading_version="1",
        needs_grading=0,
        was_suspended=1,
        status=hazards.STATUS_SUSPENDED,
        due_date="2026-08-20",
        found_at="2026-08-16 08:00:00",
        confirmed_at="2026-08-16 08:30:00",
        closed_at=None,
        created_at="2026-08-16 08:00:00",
        updated_at="2026-08-16 09:30:00",
    )
    return hazards.HazardRow(**{**base, **overrides})  # type: ignore[arg-type]


def _doc(doc_type: str, doc_no: str, **overrides: object) -> hazards.HazardDocRow:
    base: dict[str, object] = dict(
        id=1,
        hazard_no=_HAZARD_NO,
        doc_type=doc_type,
        doc_no=doc_no,
        artifact_id="a" * 32,
        photo_id=None,
        result=None,
        created_at="2026-08-16 09:30:00",
    )
    return hazards.HazardDocRow(**{**base, **overrides})  # type: ignore[arg-type]


def _ctx(**overrides: object) -> DocContext:
    base: dict[str, object] = dict(
        row=_row(),
        project_name="幸福里三期",
        signed_display="2026-08-16 09:30:00",
        due_display="8月20日(周四)",
        evidence=(
            _doc("notice", "GYT-TZ-20260816-093000-1111"),
            _doc("suspension", _SUSPENSION_NO, id=2),
        ),
    )
    return DocContext(**{**base, **overrides})  # type: ignore[arg-type]


def _reinspect_doc(**overrides: object) -> hazards.HazardDocRow:
    """一条复查记录:没有 ``artifact_id``(它不是文书),但**有自己的复查照片**。"""
    base: dict[str, object] = dict(
        doc_type="reinspect",
        doc_no=f"{_HAZARD_NO}#FC-deadbeef",
        id=3,
        artifact_id=None,
        photo_id=_REINSPECT_PHOTO_ID,
        result="fail",
        created_at="2026-08-16 10:00:00",
    )
    merged = {**base, **overrides}
    return _doc(
        str(merged.pop("doc_type")),
        str(merged.pop("doc_no")),
        **merged,
    )


def _满证据链() -> DocContext:
    """走到「上报主管部门」那一步时的证据链形态:通知单 → 暂停令 → 一条复查不合格。

    这是《监理报告》真实拿去举证的那副样子(「我通知过 + 期限到了 + 复查过 + 他没改」),
    F4 加的那一栏也只有在这副样子下才有内容。
    """
    return _ctx(
        evidence=(
            _doc("notice", "GYT-TZ-20260816-093000-1111"),
            _doc("suspension", _SUSPENSION_NO, id=2),
            _reinspect_doc(),
        )
    )


# --- 读回工具(不落盘那条钉在 test_supervision_docgen.py,这里只管内容)---------


def _paragraphs(payload: bytes) -> list[str]:
    return [p.text for p in Document(io.BytesIO(payload)).paragraphs]


def _tables(payload: bytes) -> list[list[list[str]]]:
    return [
        [[cell.text for cell in row.cells] for row in table.rows]
        for table in Document(io.BytesIO(payload)).tables
    ]


def _meta(payload: bytes) -> dict[str, str]:
    """元信息表(第一张表)读成字典。键唯一是这张表的设计前提。"""
    return {key: value for key, value in _tables(payload)[0]}


def _render(kind: DocKind, ctx: DocContext | None = None) -> bytes:
    return render_doc(kind, f"GYT-{kind.value}-20260816-093000-9999", ctx or _ctx())


def _evidence_table(payload: bytes) -> list[list[str]]:
    """证据链那张表(第二张;元信息表恒为第一张)。没有这一节就返回空列表。

    通知单与暂停令是签发链的第一步,名下本来就还没有东西,不摆这张表 —— 所以
    「没有第二张表」是正常形态,不是渲染出错。
    """
    tables = _tables(payload)
    return tables[1] if len(tables) > 1 else []


def _evidence_types(payload: bytes) -> set[str]:
    """证据链表「类型」那一列 —— 这份文书**附件真正列出来的东西**。"""
    return {行[1] for 行 in _evidence_table(payload)[1:]}


_PRIOR_CLAIM_RE = re.compile(r"已[^。;]{0,8}?(?:签发|出具)")
"""正文里「此前**已经**签发 / 已经出具」这种完成时断言的标记。

刻意要求带「已」:同一批文书里正在出的那几份,正确说法是「同批出具」——
它们在证据链里天然列不出来(证据链截止到本文书签发前,见 ``EVIDENCE_HEADING``),
所以那种说法不该被这条判据拦下。
"""

_TITLE_RE = re.compile(r"《([^》]+)》")


def _claimed_prior_docs(payload: bytes) -> set[str]:
    """正文里以完成时**点名**说"此前已经签发/出具"的文书名。

    按句切(。与;),标记与书名号必须落在**同一句**里 —— 跨句取会把
    「另行签发《工程复工令》」这种将来时也算进来。
    """
    claimed: set[str] = set()
    for 段 in _paragraphs(payload):
        for 句 in re.split(r"[。;]", 段):
            if _PRIOR_CLAIM_RE.search(句):
                claimed.update(_TITLE_RE.findall(句))
    return claimed


# ---------------------------------------------------------------------------
# 两张受控词表:漏一档不会有运行期报错,只会静默走岔
# ---------------------------------------------------------------------------


def test_五种文书每种都有正文构造函数() -> None:
    """模块导入时就会硬失败,这条用例是把那个约束写成人看得见的形式。

    漏一档的真实表现是签发时 KeyError → 兜底 500,工友只看到「系统开小差」。
    """
    from gyt.agents.supervision.documents import _SECTION_BUILDERS

    assert set(_SECTION_BUILDERS) == set(ISSUING_KINDS)


def test_五种文书的doc_type都落在hazards的受控词表里() -> None:
    """不然写库时撞 CHECK,而那要到真签发那一刻才炸。"""
    assert {k.name.lower() for k in ISSUING_KINDS} <= set(hazards.DOC_TYPES)


def test_隐患编号不是文书所以不进签发清单() -> None:
    """``DocKind.HAZARD`` 是隐患自己的身份号(doc_no.py 头注点名的两个孤儿之一),
    混进来就会去找一个并不存在的正文构造函数。"""
    assert DocKind.HAZARD not in ISSUING_KINDS


# ---------------------------------------------------------------------------
# 五种文书都要过的三关
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("kind", ISSUING_KINDS, ids=_KIND_IDS)
def test_每种文书都印着自己的编号和隐患编号(kind: DocKind) -> None:
    """编号是这份纸的对外身份 —— 电话里报它、上报主管部门时按它排证据链。
    印不上去不会有任何异常,只是这份文书从此对不回台账。
    (也正因为印在正文里,撞号重试才必须连正文一起重渲染。)
    """
    meta = _meta(_render(kind))

    assert meta[f"{DOC_TITLE_ZH[kind]}编号"] == f"GYT-{kind.value}-20260816-093000-9999"
    assert meta["隐患编号"] == _HAZARD_NO
    # 元信息表这一格印的是**首次发现**那张(hazards.photo_id)。2026-08-16 前它叫
    # 「现场照片编号」—— 与复查照片重名,而两者是完全不同的两张图(见下面那条 F4 用例)。
    assert meta[FOUND_PHOTO_LABEL] == _PHOTO_ID
    assert meta["隐患级别"] == "严重(现场判定:重大)"


@pytest.mark.parametrize("kind", ISSUING_KINDS, ids=_KIND_IDS)
def test_每种文书都盖监理那句免责且签字栏留空(kind: DocKind) -> None:
    """**五种全测**,不是只抽查一份。

    巡检记录那句讲的是「AI 初筛 vs 持证安全员」,压根没有「总监签字后生效」这一环;
    串了的表现是一份看起来很正式、却没写明"未签字不算数"的停工文书。
    签字栏冒号后面一律留空:系统替人填名字就是伪造(D15)。
    """
    texts = _paragraphs(_render(kind))

    assert SUPERVISION_DISCLAIMER in texts
    for line in SUPERVISION_SIGNATURE_LINES:
        assert line in texts


@pytest.mark.parametrize("kind", ISSUING_KINDS, ids=_KIND_IDS)
def test_每种文书的正文里都说清了是哪条隐患(kind: DocKind) -> None:
    """元信息表里有一格「隐患事项」,但正文自己也得说 —— 只在表格里出现的话,
    这份文书读起来就是一堆不知所云的要求。"""
    正文 = "\n".join(_paragraphs(_render(kind)))

    assert "临边无防护" in 正文


@pytest.mark.parametrize("kind", ISSUING_KINDS, ids=_KIND_IDS)
def test_只有监理报告的标题多一个工程监理前缀(kind: DocKind) -> None:
    """《监理报告》是报给建设主管部门的,标题得自报家门;别的四份是发给工地上的,
    标题就是它本来的名字。写反了不会报错,只是发出去的公文名字不对。"""
    标题 = _paragraphs(_render(kind))[0]
    期望 = DOC_TITLE_ZH[kind]
    if kind is DocKind.AUTHORITY_REPORT:
        期望 = f"工程监理{期望}"

    assert 标题 == 期望


# ---------------------------------------------------------------------------
# 整改期限:有就印,没有就整行不出现(不是印一个空格)
# ---------------------------------------------------------------------------


def test_有期限时元信息表印出期限那一行() -> None:
    assert _meta(_render(DocKind.NOTICE))["整改期限"] == "8月20日(周四)"


def test_没有期限时整行不出现而不是留一格空白() -> None:
    """复工令与监理报告不设新期限(``supervision_api`` 那两处传 ``due_display=""``)。
    留一格空的「整改期限:」会被读成「期限还没填」,而它本来就不该有这一行。
    """
    meta = _meta(_render(DocKind.RESUMPTION, _ctx(due_display="")))

    assert "整改期限" not in meta
    assert meta["签发时间"] == "2026-08-16 09:30:00"  # 别的行没被一起弄丢


# ---------------------------------------------------------------------------
# 证据链表:三种带表的文书靠它举证
# ---------------------------------------------------------------------------


def test_复工令正文引用原暂停令的编号() -> None:
    """复工是「原来那份暂停令作废」这件事的对应动作 —— 不写清作废的是哪一份,
    这纸复工令在证据链上就挂不住。
    """
    正文 = "\n".join(_paragraphs(_render(DocKind.RESUMPTION)))

    assert _SUSPENSION_NO in 正文


def test_证据链里没有暂停令时写未找到而不是留一段空白() -> None:
    """留空的话正文会变成「原《工程暂停令》编号 。」这种断句 —— 而它要存档。"""
    正文 = "\n".join(_paragraphs(_render(DocKind.RESUMPTION, _ctx(evidence=()))))

    assert "原《工程暂停令》编号 (未找到)。" in 正文


@pytest.mark.parametrize("kind", _带证据链的三种, ids=[k.name for k in _带证据链的三种])
def test_空证据链出一句话而不是一张空表(kind: DocKind) -> None:
    """空表在留档文件里读起来像「这一栏还没填」,而「暂无」是个明确结论 ——
    追责场合这两者意思完全不同。

    那句话还必须写清**截止到什么时候**(F5):证据链是在签发之前查的,本批正在出
    的文书它一份都列不出来。不写截止点的话,《致建设单位报告》上就是
    「正文说同批出了三份 / 附件说名下什么都没有」。
    """
    payload = _render(kind, _ctx(evidence=()))
    正文 = "\n".join(_paragraphs(payload))

    assert len(_tables(payload)) == 1  # 只剩元信息表
    assert "没有已出具的文书或复查记录" in 正文
    assert "签发前" in 正文, "空证据链那句必须写明截止点,否则读起来像「我们什么都没做」"


def test_致建设单位报告的空证据链把话接到本批三份上() -> None:
    """这一份的证据链**在生产上必然是空的**:它只从 ``open`` 那条路出,而 ``open``
    的隐患名下一份文书都还没有。通用那句「此前没有」摆在这儿,建设单位读到的就是
    「名下什么都没有」,而正文上一段刚说完同批出了三份 —— 得当场把两句接上。
    """
    正文 = "\n".join(_paragraphs(_render(DocKind.OWNER_REPORT, _ctx(evidence=()))))

    assert "同批出具的" in 正文
    assert "见「一、报告事项」" in 正文


def test_证据链按顺序列出且复查记录叫复查记录() -> None:
    """``hazard_docs.doc_type`` 是英文,印在纸上必须换成中文。

    ``reinspect`` 是唯一一个**不是文书**的档(``DocKind`` 里没有它,见 doc_no.py
    头注那两个孤儿),名字只能单独给 —— 漏了就是留档文书上印着一个 ``reinspect``。
    """
    表 = _evidence_table(_render(DocKind.AUTHORITY_REPORT, _满证据链()))

    assert 表[0] == ["序号", "类型", "编号", "结论", REINSPECT_PHOTO_LABEL, "时间"]
    assert [行[:2] for 行 in 表[1:]] == [
        ["1", "监理通知单"],
        ["2", "工程暂停令"],
        ["3", "复查记录"],
    ]
    # 没有结论的行画一个破折号,不留空格 —— 空单元格读起来像漏填
    # 结论列过反查表:库里是 pass/fail,纸上必须是中文
    # (2026-08-16 代码评审:原来这里原样印英文,而这份纸要报建设主管部门)
    assert [行[3] for 行 in 表[1:]] == ["—", "—", "不合格"]


def test_词表外的文书类型原样透出不猜() -> None:
    """将来加了一档文书却忘了同步词表:印出原字符串(看得见、查得到),
    比猜一个中文名或者当场炸掉都好 —— 这份文书本身还是要出的。
    """
    ctx = _ctx(evidence=(_doc("专项报告", "GYT-XX-20260816-093000-3333"),))

    表 = _tables(_render(DocKind.OWNER_REPORT, ctx))[1]

    assert 表[1][1] == "专项报告"


# ---------------------------------------------------------------------------
# 未归属项目(D6:空串,不是 NULL)
# ---------------------------------------------------------------------------


def test_未归属项目在纸上有明确写法而不是一格空白() -> None:
    """``supervision_api._context()`` 查不到项目时填的就是它。留白会让人以为这一格
    漏填了,写明白它就是一条待归属的隐患。"""
    assert UNASSIGNED_PROJECT_ZH.strip()

    meta = _meta(_render(DocKind.NOTICE, _ctx(project_name=UNASSIGNED_PROJECT_ZH)))

    assert meta["工程项目"] == UNASSIGNED_PROJECT_ZH


# ===========================================================================
# F4:复查照片得能从纸上取回来,而且要跟「发现时」那张分得开
# ===========================================================================


def test_两个照片编号在纸上不是同一个说法() -> None:
    """``hazards.photo_id``(首次发现)与 ``hazard_docs.photo_id``(每次复查各一张)
    是两张完全不同的照片。叫同一个名字的后果不是排版难看,是**拿发现时的照片当
    "整改后"的证据** —— 而「复查必须挂照片」这条红线的全部意义就是分得清这两者。
    """
    assert FOUND_PHOTO_LABEL != REINSPECT_PHOTO_LABEL
    # 一个不许是另一个的子串:「照片编号」/「现场照片编号」这种包含关系在纸上一样分不开,
    # 而且会让下面那条"表里只印复查那张"的断言变成假绿灯。
    assert FOUND_PHOTO_LABEL not in REINSPECT_PHOTO_LABEL
    assert REINSPECT_PHOTO_LABEL not in FOUND_PHOTO_LABEL


def test_证据链里每条复查记录印着自己的复查照片编号() -> None:
    """F4 的正主:照片编号一直存在 ``hazard_docs.photo_id`` 里,却**没有任何地方读得出来**。

    没有这一栏,「复查必须挂一张照片」就只是一道提交时的门槛:提交完谁也取不出来,
    事后追责时那张照片等于不存在 —— 而红线写下来就是为了追责。
    """
    表 = _evidence_table(_render(DocKind.AUTHORITY_REPORT, _满证据链()))
    照片列 = 表[0].index(REINSPECT_PHOTO_LABEL)

    assert [行[照片列] for 行 in 表[1:]] == ["—", "—", _REINSPECT_PHOTO_ID]
    # 印的是这一次复查那张,不是隐患首次发现那张 —— 混了就是拿发现时的照片当整改后的证据
    assert _PHOTO_ID not in {行[照片列] for 行 in 表[1:]}


def test_复查记录缺照片时纸上写得明白而不是画个破折号() -> None:
    """文书行没有照片是正常的,复查记录没有照片是**证据链断了一环**。

    两者都画成 ``—`` 的话,一条举不出现场凭据的复查记录在纸面上和一份通知单长得一样,
    而这份纸是要拿去指控施工方拒不整改的。
    """
    ctx = _ctx(evidence=(_reinspect_doc(photo_id=None),))

    表 = _evidence_table(_render(DocKind.AUTHORITY_REPORT, ctx))
    照片列 = 表[0].index(REINSPECT_PHOTO_LABEL)

    assert 表[1][照片列] == "(缺复查照片)"


def test_监理报告说清了复查照片在哪一栏且与发现时那张不是一张() -> None:
    """这份纸是报给建设主管部门的,它得让对方知道**影像调得出来、且调的是哪一张**。

    只加一栏编号不说话也行,但对方拿到的是一串三十二位十六进制 —— 不点明它是什么、
    去哪儿调,那一栏就只是噪音。
    """
    正文 = "\n".join(_paragraphs(_render(DocKind.AUTHORITY_REPORT, _满证据链())))

    assert REINSPECT_PHOTO_LABEL in 正文
    assert FOUND_PHOTO_LABEL in 正文
    assert "调取" in 正文


# ===========================================================================
# F5:正文说的话,这份文书自己的附件必须兜得住
# ===========================================================================


_证据链形态 = {
    "空证据链": (),
    "满证据链": (
        _doc("notice", "GYT-TZ-20260816-093000-1111"),
        _doc("suspension", _SUSPENSION_NO, id=2),
        _reinspect_doc(),
    ),
}
"""两副样子。**空的那一档不能省**:F5 那句 bug 原话在满证据链下是绿的
(证据链里恰好有通知单和暂停令),只有在空证据链下才现形 —— 而空证据链正是
《致建设单位报告》在生产上唯一会遇到的形态。"""


@pytest.mark.parametrize("kind", ISSUING_KINDS, ids=_KIND_IDS)
@pytest.mark.parametrize("形态", list(_证据链形态), ids=list(_证据链形态))
def test_正文说已签发的文书必须在证据链里列得出来(kind: DocKind, 形态: str) -> None:
    """🔴 **F5 的判据:一份文书的正文与它自己的附件不许互相打脸。**

    原来的《致建设单位报告》正文写死「项目监理机构已就此签发《监理通知单》与
    《工程暂停令》」,而 ``supervision_api`` 的 ``ctx = _context(...)`` 是在
    ``_sign(...)`` **之前**建的(§6.4:文件先全部落盘、库后写,顺序不能动),
    本批三份那会儿一份都还没入库 —— 于是同一张纸上,正文说签了、附表说名下什么都没有。
    这不会有任何报错,要到对方律师并排读那两段时才发现。

    判据本身:正文里凡是以完成时**点名**说"此前已签发/已出具 X"的,X 必须在这份文书
    附的证据链里列得出来。同批正在出的那几份不走这条路 —— 它们的正确说法是
    「同批出具」,证据链的截止点写在标题里(``EVIDENCE_HEADING``)。

    改这里之前先读一遍:要让正文能点名本批那几份,得把本批清单传进 ``DocContext``,
    那是 ``supervision_api._issue_documents`` 那一侧的改动。
    """
    ctx = _ctx(evidence=_证据链形态[形态])
    payload = _render(kind, ctx)

    未列出的 = _claimed_prior_docs(payload) - _evidence_types(payload)

    assert not 未列出的, (
        f"《{DOC_TITLE_ZH[kind]}》正文说此前已签发 {未列出的},"
        "而它自己附的证据链一条都列不出来 —— 正文与附件互相打脸的法律文书。"
        "本批同时出具的文书请用「同批出具」措辞:证据链截止到本文书签发前,列不出它们。"
    )


def test_致建设单位报告点名说清同批出了哪几份() -> None:
    """光把矛盾的那句删掉是不够的:建设单位最需要知道的一件事就是**工地已被要求停工**。

    所以这份报告必须点名说出同批那三份,只是措辞得是「同批出具」而不是「此前已签发」。
    """
    正文 = "\n".join(_paragraphs(_render(DocKind.OWNER_REPORT, _ctx(evidence=()))))

    for kind in OWNER_REPORT_BATCH:
        assert f"《{DOC_TITLE_ZH[kind]}》" in 正文
    assert "同批出具" in 正文
    assert "已就此签发" not in 正文, "完成时断言正是 F5 那句原话,附件兜不住它"


def test_同批清单里有报告自己且都是签发得出来的文书() -> None:
    """``OWNER_REPORT_BATCH`` 是 ``supervision_api._work_suspend`` 那个 kinds 元组的镜像
    (``render_doc`` 只拿得到自己那一份的编号,看不见同批另外两份)。

    这条只钉两件在本模块内就能验的:清单里得有报告自己(否则"本报告即其中之一"
    那句话是假的),且每一档都真的签发得出来。两边**顺序与件数**是否一致由
    ``test_supervision_api.py`` 的「一次三份且顺序固定」那条钉。
    """
    assert DocKind.OWNER_REPORT in OWNER_REPORT_BATCH
    assert set(OWNER_REPORT_BATCH) <= set(ISSUING_KINDS)


def test_证据链标题写明了截止到本文书签发前() -> None:
    """F5 的另一半:标题不写截止点,「此前无记录」与「同批出了三份」读起来就是矛盾;
    写清之后它们是互补的两句。这个标题三种带表的文书共用。"""
    assert "签发前" in EVIDENCE_HEADING

    for kind in _带证据链的三种:
        assert EVIDENCE_HEADING in _paragraphs(_render(kind, _满证据链()))


# ---------------------------------------------------------------------------
# 留档文书上不许出现英文枚举值(2026-08-16 代码评审补)
# ---------------------------------------------------------------------------


def test_整份文书里不出现pass或fail这两个英文词() -> None:
    """证据链的「结论」列原来是 ``doc.result or _NO_VALUE`` —— 把库里的枚举值
    **原样印到纸上**,而这份纸是要报建设主管部门的。

    同一份文件里 ``doc_type`` 与状态都过了反查表,唯独结论没有,口径也不一致。

    ⚠️ 断言写成「整份 docx 的文本里不含这两个词」而不是只查那一格:
    后者在换了渲染实现之后可能悄悄失效,而这条盯的是**读者真正看到的东西**。
    """
    payload = _render(DocKind.AUTHORITY_REPORT, _满证据链())
    # 段落 + 所有表格单元格 —— 结论在**表格里**,只取段落会漏掉它
    text = "\n".join(
        [*_paragraphs(payload), *(格 for 表 in _tables(payload) for 行 in 表 for 格 in 行)]
    )

    assert "不合格" in text, "复查 fail 该印成「不合格」"
    assert "fail" not in text, f"留档文书上出现了英文 fail:\n{text}"
    assert "pass" not in text, f"留档文书上出现了英文 pass:\n{text}"


def test_复查合格印成合格() -> None:
    """``pass`` 那一档单独测:上面那条走的是 fail 分支,两个取值都要有覆盖。"""
    ctx = _ctx(evidence=(_reinspect_doc(id=3, result="pass"),))
    表 = _evidence_table(_render(DocKind.AUTHORITY_REPORT, ctx))

    assert 表[1][3] == "合格"


def test_文书行没有结论时画破折号而不是留空() -> None:
    """``result is None`` 是文书行的正常形态(它本来就没有结论)。

    空单元格在留档文件里读起来像「这一栏还没填」—— 与本模块 ``_NO_VALUE`` 的原意一致。
    """
    表 = _evidence_table(_render(DocKind.AUTHORITY_REPORT, _满证据链()))

    assert 表[1][3] == "—", "通知单那一行本来就没有结论"
    assert 表[2][3] == "—", "暂停令那一行同理"
