"""supervision_api.py(监理处置直连接口)的单元测试 —— 不联网,库与产物全落 tmp_path。

环境隔离整套复用 conftest 的 ``_isolated_settings``(autouse):GYT_DATA_DIR 指到
tmp_path/data,每个用例一张全新的库、一个空的产物目录,所以本文件不需要自建环境夹具。

docx **是真渲染**(python-docx,毫秒级):这条链最容易出的错就在"渲染 → 落盘 → 写库"
的顺序与件数上,把渲染 mock 掉等于把要测的东西测没了。

这层要钉死的静默错误(全是"不报错但结果错"那一类):
  · **三条硬拦失守** —— 「一般隐患签了暂停令」与「严重隐患只发通知单」是两个方向相反的
    事故,前者平白停一片人的工,后者该停的没停;``needs_grading=1`` 放行则等于让未知
    风险按一般隐患走完闭环并被销项。三条都只在服务端代码里,提示词兜不住。
  · **``due_date`` 留空** —— 期限为空的隐患永远进不了超期清单、永远不会被升级,
    而报表上它一直显示「在办」。
  · **三份文书少出一份** —— ``documents`` 是数组,少一件就少一张下载卡,不会报错。
  · **有记录无文件** —— 库写在文件落盘之前的话,证据链里会出现一份点不开的「文书」。
  · **端点没挂进 webapp.py** —— 现象是 404,**不是启动报错**。
  · **令牌漏配时裸奔** —— langgraph 那道鉴权中间件的键漏了没有任何报错,
    只有 handler 自查这一道兜得住(与打卡链同一条上线闸)。

W10 那三条(两条 GET + 否决)另外还要钉四件,全都是"界面照常显示、只是显示错了":
  · **``project_id`` 三态塌成两态** —— 「参数不出现」被当成空串的话,「全部工地」悄悄
    变成「只有未归属」,界面上少了一大半隐患而没有任何报错(D6)。三态各一条用例。
  · **筛子静默回落** —— ``scoping.in_scope`` 认不出的词会落在「在办」且一声不吭,
    所以端点必须自己拦野词;不拦的话少给的清单在界面上看不出来少了。
  · **``filename`` 两处现拼** —— 签发时回给前端的名字与详情里回头看到的名字对不上,
    下载卡的文件名和实际落盘的就是两个东西,而不会有任何报错(用例真去比对)。
  · **复查照片取错那一张** —— ``hazard_docs.photo_id``(这次复查)与 ``hazards.photo_id``
    (首次发现)混起来 = 拿发现时的照片当"整改后"的证据,而这条红线的全部意义
    就是事后追责时分得清这两张。

⚠️ **W10 那三条端点 2026-08-16 落地时一条行为测试都没有**(那条泳道中途被打断):
路由挂载与鉴权有人盯着,函数体基本没被跑过 —— 行覆盖 87%(高于 80% 门槛、CI 照绿),
而行为覆盖是零。本文件末尾那三组(``Test隐患清单`` / ``Test隐患详情`` / ``Test否决``)
补的就是这一块。它们额外守着两条**测试自己会骗自己**的路:
  · **桩没打进去** —— 「今天」的桩必须打在 ``scoping.today_hk`` 这个**模块属性**上;
    打不进去的表现不是红,是那几条用例跟着跑测试的真实日子飘(今天绿明天红)。
    所以打完桩要先断 ``data.today`` 等于钉的那天,再去断超期。
  · **拿被告自己作证** —— 否决那几条直查 sqlite(``_库里还有``),不走 ``hazards.fetch``:
    「拒了」和「拒了但也把行删了」在回执上长得一模一样,只有库分得开。
"""

from __future__ import annotations

import re
import sqlite3
from contextlib import closing
from datetime import date
from pathlib import Path
from typing import Any, Final, NamedTuple
from urllib.parse import urlencode

import pytest
import webapp
from docx import Document as DocxDocument
from starlette.testclient import TestClient

from gyt import supervision_api
from gyt.agents.supervision import scoping
from gyt.agents.supervision import tools as supervision_tools
from gyt.agents.supervision.docgen import SUPERVISION_DISCLAIMER
from gyt.config import ALLOWED_IMAGE_EXT, get_settings
from gyt.core import artifacts
from gyt.core.artifacts import ArtifactKind
from gyt.core.doc_no import DocKind, DocNoExhaustedError, generate_unique
from gyt.db import hazards

REAL_TOKEN = "0123456789abcdef0123456789abcdef"
"""一个像样的令牌:够长、不是占位符开头 —— 两条都满足才会真正开启鉴权
(判据在 core/access.py,与 auth.py 同源)。"""

DUE_PHRASE = "下周三"
"""合法的期限原话。解析归 agents/schedule/dates.py,这里只要是它认得的写法就行。"""

FAKE_JPEG = b"\xff\xd8\xff\xe0fake-photo-bytes"

# --- 照片直传那条端点的样本字节 ------------------------------------------------
# 全是**只有文件头**的假图:上传端点按魔数判类型、不解码(它只存字节,不渲染),
# 所以文件头对了就该收下。⚠️ 这也是那条端点的能力边界 —— 它挡得住"拿别的文件冒充
# 照片",挡不住"坏图"(``_sniff_image_ext`` 的 docstring 把这笔账写明了)。

FAKE_PNG: Final[bytes] = b"\x89PNG\r\n\x1a\n" + b"fake-png-bytes"
"""PNG 的 8 字节完整签名。**别只写前四位 ``\\x89PNG``** —— 那样这条用例在
"实现只比前四位"时照样绿,而真正要钉的是它比完整签名。"""

FAKE_WEBP: Final[bytes] = b"RIFF" + b"\x24\x00\x00\x00" + b"WEBP" + b"VP8 fake"
"""WebP:``RIFF`` + 4 字节长度 + ``WEBP``。中间那四个字节是文件长度,内容随文件变 ——
所以判据必须是"第 8-12 字节等于 WEBP",不能拿一条 12 字节的前缀去比。"""

FAKE_WAV: Final[bytes] = b"RIFF" + b"\x24\x00\x00\x00" + b"WAVEfmt "
"""一段 WAV。**它和 WebP 一样以 ``RIFF`` 开头** —— 只认 ``RIFF`` 的实现会把它当照片收下。
这条样本存在的唯一意义就是钉住"两段都要比"。"""

FAKE_DOCX: Final[bytes] = b"PK\x03\x04" + b"fake-zip-bytes"
"""一份 docx(zip 容器,魔数 ``PK\\x03\\x04``)。挑它当反例是因为它是这条链上**真会出现**
的文件:监理签发的五种文书就是 docx,界面上那些下载卡随手点开一份存下来再传回复查那一栏,
是完全可能发生的误操作。"""

FAKE_TEXT: Final[bytes] = "这是一段纯文本,不是照片。".encode()
"""一段纯文本。它连魔数都没有,是"随手传了个 .txt"的样子。"""

# --- W10 那三条端点的共用常量 -------------------------------------------------
# 放在这里(而不是贴着用例)是因为 ``_new_hazard`` / ``_arrive`` 这两个搭台函数要用到。

TODAY: Final[date] = date(2026, 8, 20)
"""把「今天」钉死的那一天(周四)。

与 ``test_supervision_tools.py`` 的 ``TODAY`` **是同一天**,不是巧合:超期判据两边
共用 ``scoping.is_overdue``,同一天出发,面板的清单与对话里念的条数才好并排对照。
"""

TODAY_ISO: Final[str] = TODAY.isoformat()
YESTERDAY_ISO: Final[str] = "2026-08-19"
"""昨天到期 = 已经超期(边界严格小于,今天到期**不**算)。"""

FAR_FUTURE_ISO: Final[str] = "2026-12-31"
"""``_arrive`` 搭台时的默认期限。

**默认值必须是"远得不会超期"的那种。** 默认就超期的话,每条搭出来的隐患都自带超期,
于是那些"我压根没造超期数据"的用例会拿到超期结果 —— 而断言看起来还是绿的。
"""

PROJECT: Final[str] = "gyt-a3"
"""``_new_hazard`` 的默认工地。写成常量是因为 ``project_id`` 三态那几条要拿它做断言。"""

UNASSIGNED: Final[str] = ""
"""未归属(D6)。**是空串,不是 None** —— 三态里最容易被一个 ``or ""`` 抹平的就是它。"""

_ALL_ENDPOINTS: list[tuple[str, dict[str, Any]]] = [
    ("/supervision/confirm", {"hazard_nos": ["GYT-H-20260816-090000-0001"]}),
    ("/supervision/reject", {"hazard_no": "GYT-H-20260816-090000-0001"}),
    ("/supervision/grade", {"hazard_no": "GYT-H-20260816-090000-0001", "grade": "一般"}),
    ("/supervision/notice", {"hazard_no": "GYT-H-20260816-090000-0001", "due_phrase": DUE_PHRASE}),
    ("/supervision/suspend", {"hazard_no": "GYT-H-20260816-090000-0001", "due_phrase": DUE_PHRASE}),
    (
        "/supervision/reinspect-result",
        {"hazard_no": "GYT-H-20260816-090000-0001", "result": "pass", "after_photo_id": "0" * 32},
    ),
    ("/supervision/resume", {"hazard_no": "GYT-H-20260816-090000-0001"}),
    ("/supervision/escalate", {"hazard_no": "GYT-H-20260816-090000-0001"}),
]
"""八个 POST 端点各配一份格式合法的请求体。鉴权用例拿它逐个打 —— 鉴权必须在业务之前生效,
所以这些编号根本不存在也没关系(过了鉴权会是 404,被拦下则是 401)。"""

_GET_ENDPOINTS: list[str] = [
    "/supervision/hazards",
    "/supervision/hazards/GYT-H-20260816-090000-0001",
]
"""两条 GET 端点(W10)。**它们同样得过令牌闸** —— 隐患清单里有工地、有违规项、
有照片编号,是要登录才看得到的东西,不是公开数据;而 GET 最容易被当成"只是查一下"漏掉。"""

_ISSUING_PATHS: list[str] = [
    "/supervision/notice",
    "/supervision/suspend",
    "/supervision/resume",
    "/supervision/escalate",
]
"""四个**签发**端点。硬拦③(needs_grading=1 一律拒)必须四个全拦,漏一个就是一条
未定级的隐患从那个口子走完闭环。"""


@pytest.fixture
def client() -> TestClient:
    """真 Starlette 栈,直打本模块的 app(线上走的是 webapp.py,挂载另有用例盯着)。"""
    return TestClient(supervision_api.app)


def _photo(name: str = "整改后.jpg") -> str:
    """登记一张假照片,返回 artifact_id —— 复查结论必须挂真实存在的照片。"""
    return artifacts.register(FAKE_JPEG, kind=ArtifactKind.PHOTO, original_name=name)


def _new_hazard(
    *,
    grade: str = hazards.GRADE_NORMAL,
    severity: str = "一般",
    needs_grading: bool = False,
    item: str = "未戴安全帽",
    project_id: str = PROJECT,
) -> str:
    """登记一条隐患(落 pending),返回它的编号。编号由 core/doc_no.py 现摇,不手拼。

    ⚠️ 幂等键是 ``(project_id, photo_sha256, item)``,而 ``photo_sha256`` 是拿下面
    那几个参数拼出来的 —— **同一组参数调两次不会得到两条隐患**,``create`` 会静默
    回第一条那行(D14 的照片重传语义)。要造多条就把 ``item`` 各起各的名字
    (``test_批量确认部分成功也如实回报`` 早就是这么写的);造完顺手断一句
    「编号互不相同」,否则"我造了 4 条"其实只有 1 条,而断言看起来还是绿的。
    """
    registration = hazards.create(
        # 🔴 走 `generate_unique` 而不是裸 `new_doc_no`,**与生产代码同一条路**。
        #    编号是「到秒的时间戳 + 4 位随机」,而这些用例常常在同一秒里造几十条 ——
        #    生日悖论下撞号是必然会发生的偶发事件(2026-08-16 实测到一次:
        #    `UNIQUE constraint failed: hazards.hazard_no`,重跑就绿)。
        #    偶发红的测试比没有测试更坏:它教会人「再跑一遍就好了」,
        #    于是真回归也会被当成抖动重跑掉。
        hazard_no=generate_unique(DocKind.HAZARD, lambda no: hazards.fetch(no) is not None),
        project_id=project_id,
        photo_sha256=f"sha256-{item}-{project_id}-{grade}-{severity}-{needs_grading}",
        photo_id=_photo("现场.jpg"),
        item=item,
        severity=severity,
        grade=grade,
        grading_version="1",
        needs_grading=needs_grading,
    )
    return registration.row.hazard_no


def _arrive(status: str, *, due: str = FAR_FUTURE_ISO, **kwargs: Any) -> str:
    """把一条隐患**直接用 db 层**推到指定状态,返回编号。

    刻意绕开端点:端点带着三条硬拦,而这里要搭的台子恰恰包括"硬拦本该拦下的局面"
    (比如一条 ``needs_grading=1`` 却已经停过工的隐患 —— 从端点根本走不到)。
    db 层只管状态合法性、不管业务判断,正好当搭台工具。

    ``due`` 只对会写期限的那几档有意义(notified / suspended,以及从它们派生出来的
    reinspect_failed / resuming / closed / escalated)。它是 W10 的清单用例加的:
    超期与"今天到期"这两档只能靠期限本身造出来,而**默认值必须是不超期的**
    (理由在 ``FAR_FUTURE_ISO``)。

    ``closed`` / ``escalated`` 两档同样是 W10 加的:「全部」这个筛子与「在办」的**唯一**
    差别就是这两档在不在(``scoping.CLOSED_STATUSES``)。造数里没有它们的话,
    「全部一条都不筛」那条用例测了个寂寞 —— 两个筛子返回同一批,断言照样绿。
    """
    hazard_no = _new_hazard(**kwargs)
    if status == hazards.STATUS_PENDING:
        return hazard_no
    hazards.confirm(hazard_no)
    if status == hazards.STATUS_OPEN:
        return hazard_no
    if status == hazards.STATUS_NOTIFIED:
        hazards.mark_notified(hazard_no, due)
        return hazard_no
    if status == hazards.STATUS_SUSPENDED:
        hazards.mark_suspended(hazard_no, due)
        return hazard_no
    if status == hazards.STATUS_REINSPECT_FAILED:
        hazards.mark_notified(hazard_no, due)
        hazards.mark_reinspect_failed(hazard_no)
        return hazard_no
    if status == hazards.STATUS_RESUMING:
        hazards.mark_suspended(hazard_no, due)
        hazards.pass_reinspection(hazard_no)
        return hazard_no
    if status == hazards.STATUS_CLOSED:
        # 没停过工的复查合格 → 直接 closed(挑边由 db 按 was_suspended 决定)
        hazards.mark_notified(hazard_no, due)
        hazards.pass_reinspection(hazard_no)
        return hazard_no
    if status == hazards.STATUS_ESCALATED:
        hazards.mark_notified(hazard_no, due)
        hazards.mark_reinspect_failed(hazard_no)
        hazards.mark_escalated(hazard_no)
        return hazard_no
    raise AssertionError(f"搭台脚本还不会造 {status} 这个状态")  # pragma: no cover


def _docx_files() -> list[Path]:
    """产物目录里所有落了盘的 docx —— 用来分辨"文件出了没"与"库写了没"。"""
    return sorted(get_settings().artifacts_dir.glob("*/*.docx"))


def _产物文件() -> set[Path]:
    """产物目录里现有的**全部**文件(正文 blob + sidecar)。

    上传那条端点的几条"被拒"用例拿它前后对比:被拒的请求**一个字节都不许落盘**。
    只断状态码是不够的 —— 「400 拒了」和「400 拒了但也把那份垃圾存下来了」在回执上
    长得一模一样,而后者会让产物目录被随手一传就撑大,且没有任何报错。

    ⚠️ 断"集合相等"而不是"数量没变":同一秒里另一份文件恰好被删又被建的话,
    数量能对上而内容已经变了(这条链上真会同时有几份产物在动)。
    """
    return {p for p in get_settings().artifacts_dir.rglob("*") if p.is_file()}


# ---------------------------------------------------------------------------
# 挂载与词表:漏了都不报错
# ---------------------------------------------------------------------------


def test_十一条路由都挂进了webapp() -> None:
    """webapp.py 是自定义路由唯一的挂载点,漏铺 = 全部 404 且**没有任何启动报错**。

    ⚠️ 2026-08-16(W10·S4)实测到这条用例守不住的那一半,别把它当成全部保障:
    它验的是「``SUPERVISION_ROUTES`` 里的路由都在 ``webapp.app.routes`` 里」,
    而线上/本机容器还要求 **``webapp.py`` 这个文件本身真的在容器里、
    ``langgraph.json`` 里真的有 http 块**。当时本机镜像两样都缺(镜像比
    webapp.py 早四天),于是端点全 404 —— **而这条用例照样是绿的**,
    因为它在宿主的 Python 进程里直接 import,根本不经过容器。
    容器那一侧靠 ``docker-compose.dev.yml`` 把三份 src/ 外的源码挂进去来保证。
    """
    mounted = {route.path for route in webapp.app.routes}
    declared = {route.path for route in supervision_api.SUPERVISION_ROUTES}

    assert declared <= mounted, f"这些路由没挂进 webapp.py:{sorted(declared - mounted)}"
    # 八条 POST(W9 七条 + W10 的 reject)+ 两条 GET(W10)+ 一条上传(照片直传)
    assert len(declared) == 11, "加了路由要连同这个数一起改 —— 它是「有没有漏铺」的对账锚"
    # 逐条点名 W10 那三条与上传那条:只对总数会在「删一条旧的、加一条新的」时对上而失效。
    for 新路径 in (
        "/supervision/hazards",
        "/supervision/hazards/{hazard_no}",
        "/supervision/reject",
        "/supervision/photo",
    ):
        assert 新路径 in declared and 新路径 in mounted, f"{新路径} 没挂上"


def test_八个状态的中文名取自scoping而不是本模块另抄一份() -> None:
    """状态词表漂了的表现是回给工友的话里冒出一个英文词。导入期已有硬失败守卫,
    这条用例是把那个约束写成人看得见的形式。

    🔴 2026-08-16(W10·S1/S2)之前,本模块有一份 ``_STATUS_ZH`` 与
    ``agents/supervision/tools.py`` 那份**逐字节相同**的拷贝(当时用 AST 比对过,
    连内嵌注释都一样)。现在两份都收敛到了 ``agents/supervision/scoping.py``。
    下面第二条断言守的就是**别再抄回来**:哪天有人图省事在本模块里重建一份,
    两份会各自漂移,而两边测试都绿。
    """
    assert set(scoping.STATUS_ZH) == set(hazards.STATUSES)
    assert not hasattr(supervision_api, "_STATUS_ZH"), (
        "本模块不许再有自己的状态中文名词表 —— 唯一真相在 agents/supervision/scoping.py"
    )


def test_文书类型的中文名两份拷贝逐字相同() -> None:
    """``supervision_api._DOC_TYPE_ZH`` 与 ``agents/supervision/tools.py`` 那份是**两份**。

    暂时收敛不掉:tools.py 会拉起 langchain,而本模块是被 langgraph 按**文件路径**加载的
    HTTP 层;下沉到 ``scoping.py`` 又会破它那张被 AST 钉死的 import 白名单
    (那边没有 ``core/doc_no``)。五档文书名两边都是从 ``core/doc_no.DOC_TITLE_ZH``
    **派生**的,唯一手写的格子是「reinspect → 复查记录」那一行。

    漂开的表现:同一份文书在操作台上和在对话里叫两个名字,而**两边测试各自都绿** ——
    各自模块内那道导入期守卫只验"每个 doc_type 都有名字",验不了"两份名字一样"。

    🔴 这条 2026-08-16 才补上。在那之前 ``supervision_api._DOC_TYPE_ZH`` 的 docstring
    已经写着"``test_supervision_api.py`` 有一条断言把两份钉成逐字相同",而那是张空头支票:
    实际存在的只有文件头上一个没人用的 ``supervision_tools`` import(ruff F401 报着)。
    """
    assert supervision_api._DOC_TYPE_ZH == supervision_tools._DOC_TYPE_ZH


# ---------------------------------------------------------------------------
# 鉴权(纵深防御第二道):判定与 auth.py 同源
# ---------------------------------------------------------------------------


class Test鉴权:
    @pytest.mark.parametrize(("path", "body"), _ALL_ENDPOINTS)
    def test_设了令牌但没带钥匙一律拒(
        self, client: TestClient, monkeypatch: pytest.MonkeyPatch, path: str, body: dict[str, Any]
    ) -> None:
        """八个 POST 端点全测:漏一个就是一条能绕过令牌的写入路径。"""
        monkeypatch.setenv("GYT_ACCESS_TOKEN", REAL_TOKEN)
        get_settings.cache_clear()

        resp = client.post(path, json=body)

        assert resp.status_code == 401
        assert resp.json()["error_code"] == "UNAUTHORIZED"

    @pytest.mark.parametrize("path", _GET_ENDPOINTS)
    def test_两条GET也得过令牌闸(
        self, client: TestClient, monkeypatch: pytest.MonkeyPatch, path: str
    ) -> None:
        """W10 的两条查询端点同样要拦。

        单列一条而不是并进上面那条,是因为**它们没有请求体**,``client.post(json=...)``
        套不上;而合并成一个「路径 + 方法」的大表会让上面那条读起来像是 GET 也在测。
        真正的风险在于:GET 最容易被当成"只是查一下"而漏掉鉴权,可隐患清单里有工地名、
        违规项、照片编号 —— 那是登录之后才该看到的东西。
        """
        monkeypatch.setenv("GYT_ACCESS_TOKEN", REAL_TOKEN)
        get_settings.cache_clear()

        resp = client.get(path)

        assert resp.status_code == 401
        assert resp.json()["error_code"] == "UNAUTHORIZED"

    def test_上传照片也得过令牌闸(
        self, client: TestClient, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """🔴 上传口不拦 = 给任何路过的人一个**往这台机器写文件**的入口。

        单列一条而不是并进上面那张表:它的请求体是**原始字节**,``client.post(json=...)``
        套不上(那正是它要另开一个外壳的原因)。

        断言里那句「盘上一个字节都没多」是关键的另一半:401 与「401 但已经把 10MB
        收进来落了盘」在回执上长得一模一样,只有产物目录分得开 ——
        ``_serve`` 里令牌自查排在读 body 之前,钉的就是这个顺序。
        """
        monkeypatch.setenv("GYT_ACCESS_TOKEN", REAL_TOKEN)
        get_settings.cache_clear()
        之前 = _产物文件()

        resp = client.post("/supervision/photo", content=FAKE_JPEG)

        assert resp.status_code == 401
        assert resp.json()["error_code"] == "UNAUTHORIZED"
        assert _产物文件() == 之前, "被拒的请求不许在盘上留下任何东西"

    def test_GET带对令牌就放行到业务层(
        self, client: TestClient, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """放行的证据:清单端点开始回 200(空台账也是 ok),而不是 401。"""
        monkeypatch.setenv("GYT_ACCESS_TOKEN", REAL_TOKEN)
        get_settings.cache_clear()

        resp = client.get("/supervision/hazards", headers={"X-Api-Key": REAL_TOKEN})

        assert resp.status_code == 200
        assert resp.json()["ok"] is True

    def test_带对令牌就放行到业务层(
        self, client: TestClient, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """放行的证据是它开始回业务错误(编号查不到 = 404),而不是 401。"""
        monkeypatch.setenv("GYT_ACCESS_TOKEN", REAL_TOKEN)
        get_settings.cache_clear()

        resp = client.post(
            "/supervision/notice",
            json={"hazard_no": "GYT-H-20260816-090000-0001", "due_phrase": DUE_PHRASE},
            headers={"X-Api-Key": REAL_TOKEN},
        )

        assert resp.status_code == 404
        assert resp.json()["error_code"] == "NOT_FOUND"

    def test_没配令牌时不拦(self, client: TestClient) -> None:
        """本机联调与真机验收都不发令牌头,「未配置也拦」等于当场打死它们。"""
        resp = client.post("/supervision/resume", json={"hazard_no": "查无此号"})

        assert resp.status_code == 404  # 走到业务层了


# ---------------------------------------------------------------------------
# 🔴 三条硬拦
# ---------------------------------------------------------------------------


class Test三条硬拦:
    def test_严重隐患不许只签通知单(self, client: TestClient) -> None:
        """硬拦①:该走三文书那条路的,不许从通知单这个口子溜出去。"""
        hazard_no = _arrive(hazards.STATUS_OPEN, grade=hazards.GRADE_SEVERE, severity="重大")

        resp = client.post(
            "/supervision/notice", json={"hazard_no": hazard_no, "due_phrase": DUE_PHRASE}
        )

        assert resp.status_code == 409
        body = resp.json()
        assert body["error_code"] == "CONFLICT"
        assert "暂停令" in body["user_msg"]  # 得告诉人去哪条路,不是光说不行
        # 状态一动没动,盘上也没有半份文书
        assert hazards.fetch(hazard_no).status == hazards.STATUS_OPEN
        assert _docx_files() == []

    def test_一般隐患不许签暂停令(self, client: TestClient) -> None:
        """硬拦②:反方向那条 —— 签下去就是平白停一片人的工。"""
        hazard_no = _arrive(hazards.STATUS_OPEN)

        resp = client.post(
            "/supervision/suspend", json={"hazard_no": hazard_no, "due_phrase": DUE_PHRASE}
        )

        assert resp.status_code == 409
        assert "通知单" in resp.json()["user_msg"]
        assert hazards.fetch(hazard_no).status == hazards.STATUS_OPEN
        assert _docx_files() == []

    @pytest.mark.parametrize("path", _ISSUING_PATHS)
    def test_未定级的隐患任何签发都被拒(self, client: TestClient, path: str) -> None:
        """硬拦③(Codex#11):四个签发端点全拦。

        各自搭到那个端点该有的状态(用 db 层直推,见 ``_arrive`` 的说明),
        这样被拒的原因只可能是这面旗子,不是状态不对。
        """
        status_for = {
            "/supervision/notice": hazards.STATUS_OPEN,
            "/supervision/suspend": hazards.STATUS_OPEN,
            "/supervision/resume": hazards.STATUS_RESUMING,
            "/supervision/escalate": hazards.STATUS_REINSPECT_FAILED,
        }
        grade = hazards.GRADE_SEVERE if path == "/supervision/suspend" else hazards.GRADE_NORMAL
        hazard_no = _arrive(status_for[path], grade=grade, severity="待定级", needs_grading=True)
        before = hazards.fetch(hazard_no).status

        resp = client.post(path, json={"hazard_no": hazard_no, "due_phrase": DUE_PHRASE})

        assert resp.status_code == 409
        assert "定级" in resp.json()["user_msg"]
        assert hazards.fetch(hazard_no).status == before
        assert _docx_files() == []


# ---------------------------------------------------------------------------
# 整改期限:解析不出就 fail,绝不留空
# ---------------------------------------------------------------------------


class Test整改期限:
    def test_期限看不懂就拒且什么都没发生(self, client: TestClient) -> None:
        """期限留空的隐患永远进不了超期清单,也就永远不会被升级 —— 所以宁可这次不签。"""
        hazard_no = _arrive(hazards.STATUS_OPEN)

        resp = client.post(
            "/supervision/notice", json={"hazard_no": hazard_no, "due_phrase": "改快一点"}
        )

        assert resp.status_code == 400
        assert resp.json()["error_code"] == "INVALID_INPUT"
        # dates.py 的报错本身就是人话,还附了正确写法 —— 原样透出去比重新包一层有用
        assert "没看懂" in resp.json()["user_msg"]
        assert hazards.fetch(hazard_no).due_date is None
        assert hazards.fetch(hazard_no).status == hazards.STATUS_OPEN
        assert _docx_files() == []

    def test_不给期限也拒(self, client: TestClient) -> None:
        """空 due_phrase 不许被当成「无期限」放过去(schedule 的任务可以没期限,隐患不行)。"""
        hazard_no = _arrive(hazards.STATUS_OPEN)

        resp = client.post("/supervision/notice", json={"hazard_no": hazard_no})

        assert resp.status_code == 400
        assert "期限" in resp.json()["user_msg"]
        assert hazards.fetch(hazard_no).status == hazards.STATUS_OPEN

    def test_期限按香港日历日写进台账(self, client: TestClient) -> None:
        """签发成功时 due_date 必须是解析出来的真日期(ISO),不是原话。"""
        hazard_no = _arrive(hazards.STATUS_OPEN)

        resp = client.post(
            "/supervision/notice", json={"hazard_no": hazard_no, "due_phrase": "明天"}
        )

        assert resp.status_code == 200
        due = hazards.fetch(hazard_no).due_date
        assert due is not None
        assert len(due) == len("2026-08-17") and due.count("-") == 2


# ---------------------------------------------------------------------------
# 🔴 暂停令:三份文书、顺序固定、文件先落盘库后写
# ---------------------------------------------------------------------------


class Test暂停令三文书:
    def _suspend(self, client: TestClient) -> tuple[str, dict[str, Any]]:
        hazard_no = _arrive(hazards.STATUS_OPEN, grade=hazards.GRADE_SEVERE, severity="重大")
        resp = client.post(
            "/supervision/suspend", json={"hazard_no": hazard_no, "due_phrase": DUE_PHRASE}
        )
        assert resp.status_code == 200, resp.text
        return hazard_no, resp.json()

    def test_一次三份且顺序固定(self, client: TestClient) -> None:
        """顺序是冻结契约的一部分:前端按数组渲染 N 张下载卡,乱序 = 卡片对不上人的预期。"""
        _, body = self._suspend(client)

        docs = body["data"]["documents"]
        assert [d["doc_type"] for d in docs] == ["notice", "suspension", "owner_report"]

    def test_三份的种类与顺序与文书侧那份镜像逐项相同(self, client: TestClient) -> None:
        """`documents.OWNER_REPORT_BATCH` 是本端点 kinds 元组的**跨文件镜像**。

        为什么它必须存在:《致建设单位报告》的正文要点名「同批出具了哪几份」,
        而渲染发生在落库**之前**(§6.4 冻结的顺序),它看不见同批另外两份的编号 ——
        只能靠一份写死的清单。清单与真实签发漂开的后果是**正文说签了三份、
        实际出了两份**,而那是要送到建设单位手里的纸。

        两边各自的模块内守卫只验「清单里得有报告自己、都是签发得出来的文书」,
        **验不了顺序与件数** —— 那正是这条断言的位置(2026-08-16 代码评审补)。
        """
        from gyt.agents.supervision.documents import OWNER_REPORT_BATCH

        _, body = self._suspend(client)

        assert [d["doc_type"] for d in body["data"]["documents"]] == [
            kind.name.lower() for kind in OWNER_REPORT_BATCH
        ]

    def test_每一项四个键齐全且编号类型段对得上(self, client: TestClient) -> None:
        """少一个键前端就少渲一块,而且不会报错。编号的类型段是这份文书的对外身份。"""
        _, body = self._suspend(client)

        prefixes = {"notice": "GYT-TZ-", "suspension": "GYT-ZT-", "owner_report": "GYT-JS-"}
        for doc in body["data"]["documents"]:
            assert set(doc) == {"doc_type", "doc_no", "artifact_id", "filename"}
            assert doc["doc_no"].startswith(prefixes[doc["doc_type"]])
            assert doc["filename"].endswith(".docx")

    def test_三份都真落了盘也真进了证据链(self, client: TestClient) -> None:
        """「有记录无文件」是这条链最危险的失效:证据链里一份点不开的文书。"""
        hazard_no, body = self._suspend(client)

        for doc in body["data"]["documents"]:
            assert artifacts.resolve(doc["artifact_id"]).is_file()
        rows = hazards.docs_of([hazard_no])
        assert [r.doc_type for r in rows] == ["notice", "suspension", "owner_report"]
        assert all(r.artifact_id for r in rows)
        assert len(_docx_files()) == 3

    def test_状态与期限一起落库(self, client: TestClient) -> None:
        hazard_no, body = self._suspend(client)

        row = hazards.fetch(hazard_no)
        assert row.status == hazards.STATUS_SUSPENDED == body["data"]["status"]
        assert row.due_date is not None
        assert row.was_suspended == 1  # 此后复查合格也必须先出复工令

    def test_三份共用一个时间快照(self, client: TestClient) -> None:
        """编号里的日期时刻必须一致 —— 各取各的 now,跨秒时同一次签发在纸面上会变成
        三个时间点,而追责时先被质疑的就是文件本身(doc_no.py 头注的原话)。"""
        _, body = self._suspend(client)

        # GYT-TZ-20260816-093000-abcd → 取「20260816-093000」那一段
        时刻 = {"-".join(d["doc_no"].split("-")[2:4]) for d in body["data"]["documents"]}
        assert len(时刻) == 1

    def test_文书正文印着编号且用的是监理那句免责(self, client: TestClient) -> None:
        """两件事一起盯:

        · **编号必须印在纸上** —— 它是这份文书的对外身份,也正因为印在正文里,
          撞号重试才必须连正文一起重渲染;
        · **免责句不许串** —— 巡检记录那句讲的是「AI 初筛 vs 持证安全员」,压根没有
          「总监签字后生效」这一环。串了的表现是一份看起来很正式、却没写明"未签字不算数"
          的停工文书(docgen.py 头注点名的那条)。
        """
        hazard_no, body = self._suspend(client)
        暂停令 = body["data"]["documents"][1]

        parsed = DocxDocument(str(artifacts.resolve(暂停令["artifact_id"])))
        正文 = "\n".join(p.text for p in parsed.paragraphs)
        表格 = [c.text for t in parsed.tables for r in t.rows for c in r.cells]

        assert parsed.paragraphs[0].text == "工程暂停令"
        assert 暂停令["doc_no"] in 表格
        assert hazard_no in 表格
        assert SUPERVISION_DISCLAIMER in 正文
        # 签字栏留空:系统替人填名字就是伪造(D15)
        assert "总监理工程师(签字):" in 正文

    def test_措辞是已出具暂停令而不是已停工(self, client: TestClient) -> None:
        """suspended 只证明文书出了稿,不证明工地真停了工(db 的 STATUSES 头注)。
        对外措辞升级一格,就是拿一份没签字的稿子去宣称已经责令停工。"""
        _, body = self._suspend(client)

        user_msg = body["user_msg"]
        assert "出稿" in user_msg
        assert "签字" in user_msg
        assert "已责令停工" not in user_msg

    def test_库写失败时留孤儿文件但状态一动没动(
        self, client: TestClient, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """§6.4 的取舍:孤儿文件(有文件无记录)可以接受,孤儿记录绝对不行。

        文件先落盘,所以库炸的时候盘上已经有三份 —— 这正是预期,不是 bug。
        真正要断的是**状态没变**:库那一步是一个事务,迁移没成就一行都没写。
        """
        hazard_no = _arrive(hazards.STATUS_OPEN, grade=hazards.GRADE_SEVERE, severity="重大")

        def 炸(*_args: Any, **_kwargs: Any) -> bool:
            raise RuntimeError("台账挂了")

        monkeypatch.setattr(hazards, "mark_suspended", 炸)

        resp = client.post(
            "/supervision/suspend", json={"hazard_no": hazard_no, "due_phrase": DUE_PHRASE}
        )

        assert resp.status_code == 500
        assert resp.json()["error_code"] == "INTERNAL"
        assert hazards.fetch(hazard_no).status == hazards.STATUS_OPEN
        assert hazards.docs_of([hazard_no]) == []
        assert len(_docx_files()) == 3  # 孤儿,交清理器


# ---------------------------------------------------------------------------
# 编号:撞库重试与"摇不出来"
# ---------------------------------------------------------------------------


class Test文书编号:
    def test_撞库时整轮重摇号重渲染(
        self, client: TestClient, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """编号印在文书正文里,所以撞号必须连正文一起重来 —— 只换库里那一列的话,
        纸上的号和台账对不上,那份文书自己证伪自己。"""
        hazard_no = _arrive(hazards.STATUS_OPEN)
        真的 = hazards.mark_notified
        调用次数 = {"n": 0}

        def 先撞一次(*args: Any, **kwargs: Any) -> bool:
            调用次数["n"] += 1
            if 调用次数["n"] == 1:
                raise sqlite3.IntegrityError("UNIQUE constraint failed: hazard_docs.doc_no")
            return 真的(*args, **kwargs)

        monkeypatch.setattr(hazards, "mark_notified", 先撞一次)

        resp = client.post(
            "/supervision/notice", json={"hazard_no": hazard_no, "due_phrase": DUE_PHRASE}
        )

        assert resp.status_code == 200
        assert 调用次数["n"] == 2
        # 库里那份与回给前端的那份是同一个号(重渲染之后两边仍然对得上)
        回执号 = resp.json()["data"]["documents"][0]["doc_no"]
        assert [r.doc_no for r in hazards.docs_of([hazard_no])] == [回执号]
        assert len(_docx_files()) == 2  # 第一轮那份成了孤儿

    def test_摇不出编号时接住并说人话(
        self, client: TestClient, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """**不许降级成「用最后那个撞了的号继续」** —— 两份文书顶着同一个编号,
        比这次没出成严重得多。异常原文只进日志。"""
        hazard_no = _arrive(hazards.STATUS_OPEN)

        def 摇不出来(*_args: Any, **_kwargs: Any) -> str:
            raise DocNoExhaustedError("连着摇了 5 次监理通知单编号都已经被占用")

        monkeypatch.setattr(supervision_api, "generate_unique", 摇不出来)

        resp = client.post(
            "/supervision/notice", json={"hazard_no": hazard_no, "due_phrase": DUE_PHRASE}
        )

        assert resp.status_code == 500
        body = resp.json()
        assert body["error_code"] == "INTERNAL"
        assert "没有出" in body["user_msg"]  # 明确告诉工友:什么都没发生
        assert "DocNoExhausted" not in body["user_msg"]  # 内部细节不许露头
        assert hazards.fetch(hazard_no).status == hazards.STATUS_OPEN
        assert _docx_files() == []


# ---------------------------------------------------------------------------
# 状态机:非法迁移一律拒,而且要说清现在是哪一站
# ---------------------------------------------------------------------------


class Test状态机:
    def test_待确认的隐患不能直接签通知单(self, client: TestClient) -> None:
        """没有这道闸,safety 看一眼照片就能开启法律流程(D17 的存在理由)。"""
        hazard_no = _arrive(hazards.STATUS_PENDING)

        resp = client.post(
            "/supervision/notice", json={"hazard_no": hazard_no, "due_phrase": DUE_PHRASE}
        )

        assert resp.status_code == 409
        assert "待确认" in resp.json()["user_msg"]
        assert hazards.fetch(hazard_no).status == hazards.STATUS_PENDING

    def test_签过通知单的不能再签一次(self, client: TestClient) -> None:
        hazard_no = _arrive(hazards.STATUS_NOTIFIED)

        resp = client.post(
            "/supervision/notice", json={"hazard_no": hazard_no, "due_phrase": DUE_PHRASE}
        )

        assert resp.status_code == 409
        assert _docx_files() == []

    def test_没停过工的不许发复工令(self, client: TestClient) -> None:
        """滥发复工令:一条从没停过工的隐患拿到一纸复工令,等于凭空承认它停过。"""
        hazard_no = _arrive(hazards.STATUS_NOTIFIED)

        resp = client.post("/supervision/resume", json={"hazard_no": hazard_no})

        assert resp.status_code == 409
        assert "已签发通知单" in resp.json()["user_msg"]

    def test_没复查过不许上报主管部门(self, client: TestClient) -> None:
        """升级的举证链是「我通知过 + 期限到了 + 复查过 + 他没改」,少一环就是拿
        建立在自己记乱账上的材料去指控施工方。"""
        hazard_no = _arrive(hazards.STATUS_NOTIFIED)

        resp = client.post("/supervision/escalate", json={"hazard_no": hazard_no})

        assert resp.status_code == 409
        assert hazards.fetch(hazard_no).status == hazards.STATUS_NOTIFIED

    def test_查无此隐患是404(self, client: TestClient) -> None:
        resp = client.post(
            "/supervision/notice",
            json={"hazard_no": "GYT-H-20260101-000000-ffff", "due_phrase": "明天"},
        )

        assert resp.status_code == 404
        assert resp.json()["error_code"] == "NOT_FOUND"


# ---------------------------------------------------------------------------
# 批量确认(D17 的人工闸)
# ---------------------------------------------------------------------------


class Test批量确认:
    def test_批量确认部分成功也如实回报(self, client: TestClient) -> None:
        """一条失败不牵连其它条(db 层要求每条一个事务),但失败的必须**出现在回执里** ——
        静默吞掉的话,监理以为都确认了,少的那条从此没人管。"""
        好的 = [_arrive(hazards.STATUS_PENDING, item=f"隐患{i}") for i in range(2)]
        已确认 = _arrive(hazards.STATUS_OPEN, item="已经确认过的")

        resp = client.post(
            "/supervision/confirm",
            json={"hazard_nos": [*好的, 已确认, "GYT-H-20260101-000000-ffff"]},
        )

        assert resp.status_code == 200
        data = resp.json()["data"]
        assert data["confirmed"] == 好的
        assert [f["hazard_no"] for f in data["failed"]] == [
            已确认,
            "GYT-H-20260101-000000-ffff",
        ]
        assert all(hazards.fetch(no).status == hazards.STATUS_OPEN for no in 好的)

    def test_单条也收(self, client: TestClient) -> None:
        hazard_no = _arrive(hazards.STATUS_PENDING)

        resp = client.post("/supervision/confirm", json={"hazard_no": hazard_no})

        assert resp.status_code == 200
        assert resp.json()["data"]["confirmed"] == [hazard_no]

    def test_重复编号只算一条(self, client: TestClient) -> None:
        """界面上双击会把同一条发两遍,不去重的话第二条被报成「不用再确认」,
        看起来像出了错,其实什么问题都没有。"""
        hazard_no = _arrive(hazards.STATUS_PENDING)

        resp = client.post("/supervision/confirm", json={"hazard_nos": [hazard_no, hazard_no]})

        assert resp.status_code == 200
        data = resp.json()["data"]
        assert data["confirmed"] == [hazard_no]
        assert data["failed"] == []

    def test_一条都没确认成回409(self, client: TestClient) -> None:
        resp = client.post("/supervision/confirm", json={"hazard_nos": ["查无此号"]})

        assert resp.status_code == 409
        assert "查不到" in resp.json()["user_msg"]

    def test_空清单被拒(self, client: TestClient) -> None:
        resp = client.post("/supervision/confirm", json={"hazard_nos": []})

        assert resp.status_code == 400


# ---------------------------------------------------------------------------
# 人工定级:未定级隐患唯一的解锁通道
# ---------------------------------------------------------------------------


class Test人工定级:
    def test_定级之后旗子清掉且签发解锁(self, client: TestClient) -> None:
        hazard_no = _arrive(hazards.STATUS_OPEN, severity="待定级", needs_grading=True)

        graded = client.post(
            "/supervision/grade", json={"hazard_no": hazard_no, "grade": hazards.GRADE_NORMAL}
        )
        signed = client.post(
            "/supervision/notice", json={"hazard_no": hazard_no, "due_phrase": DUE_PHRASE}
        )

        assert graded.status_code == 200
        assert graded.json()["data"]["needs_grading"] is False
        assert hazards.fetch(hazard_no).needs_grading == 0
        assert signed.status_code == 200  # 解锁了

    def test_级别不在词表被拒(self, client: TestClient) -> None:
        hazard_no = _arrive(hazards.STATUS_OPEN)

        resp = client.post("/supervision/grade", json={"hazard_no": hazard_no, "grade": "特别严重"})

        assert resp.status_code == 400
        assert hazards.fetch(hazard_no).grade == hazards.GRADE_NORMAL

    def test_签过文书之后不许改级(self, client: TestClient) -> None:
        """改了的话,已经发出去的那份文书和台账当场对不上 —— 而本批不做重签。"""
        hazard_no = _arrive(hazards.STATUS_NOTIFIED)

        resp = client.post(
            "/supervision/grade", json={"hazard_no": hazard_no, "grade": hazards.GRADE_SEVERE}
        )

        assert resp.status_code == 409
        assert hazards.fetch(hazard_no).grade == hazards.GRADE_NORMAL

    def test_可改的状态名单从db那一份派生(self) -> None:
        """两份名单漂开的表现是端点先放行、库再拒,工友拿到「状态刚被改过」这种
        驴唇不对马嘴的提示 —— 而两边看各自的代码都觉得自己没错。"""
        assert supervision_api._GRADABLE_STATUSES == frozenset(hazards.GRADABLE_STATUSES)

    def test_先读之后被抢签了文书_定级必须落空(
        self, client: TestClient, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """🔴 并发窗口:端点是「先读 status 判断能不能改 → 再 UPDATE」两步。

        这里不真起两个线程,手工在两步之间把库改掉 —— 效果一样,而且必然复现:
        读到 ``open``(可以改级)之后、UPDATE 之前,另一个请求把《监理通知单》按
        「一般隐患」签了(状态成 notified、文书已落盘)。这时定级要是照样命中,
        库里就成了「严重隐患 + 已按一般隐患发出去的通知单」,**而且一声不吭** ——
        该停的工没停,而台账看起来一切正常。

        兜得住这一下的**只有 db 那条 ``UPDATE … WHERE status IN (…)``**:
        端点先读的那份快照说可以,它说了不算。
        """
        hazard_no = _arrive(hazards.STATUS_OPEN, severity="待定级", needs_grading=True)
        真的 = hazards.set_grade

        def 抢在前面(*args: Any, **kwargs: Any) -> bool:
            # 「另一个请求」:先读之后、这条 UPDATE 之前把通知单签了
            hazards.mark_notified(
                hazard_no,
                "2026-12-31",
                docs=[hazards.DocDraft("notice", "GYT-TZ-抢跑", artifact_id="a" * 32)],
            )
            return 真的(*args, **kwargs)

        monkeypatch.setattr(hazards, "set_grade", 抢在前面)

        resp = client.post(
            "/supervision/grade", json={"hazard_no": hazard_no, "grade": hazards.GRADE_SEVERE}
        )

        assert resp.status_code == 409
        assert resp.json()["error_code"] == "CONFLICT"
        assert "刚被改过" in resp.json()["user_msg"]  # 人话:让人刷新再看,而不是报个内部词
        row = hazards.fetch(hazard_no)
        # 两个见证:级别没被改成严重,待定级的旗子也没被顺手清掉(清了签发闸就白开了)
        assert row.grade == hazards.GRADE_NORMAL
        assert row.needs_grading == 1
        assert row.status == hazards.STATUS_NOTIFIED  # 抢跑那一手确实落了库


# ---------------------------------------------------------------------------
# 复查结论:照片是证据,结论由人下,落到哪一站由 db 挑边
# ---------------------------------------------------------------------------


class Test复查结论:
    def test_缺复查照片必拒(self, client: TestClient) -> None:
        """拿不到照片就没有任何路径能把隐患改成 closed(方案 §5.2 红线)。"""
        hazard_no = _arrive(hazards.STATUS_NOTIFIED)

        resp = client.post(
            "/supervision/reinspect-result", json={"hazard_no": hazard_no, "result": "pass"}
        )

        assert resp.status_code == 400
        assert "照片" in resp.json()["user_msg"]
        assert hazards.fetch(hazard_no).status == hazards.STATUS_NOTIFIED

    def test_照片编号打错也拒(self, client: TestClient) -> None:
        """只校验"非空"是不够的:编号错一位照样非空,而复查合格是能把隐患销项的。"""
        hazard_no = _arrive(hazards.STATUS_NOTIFIED)

        resp = client.post(
            "/supervision/reinspect-result",
            json={"hazard_no": hazard_no, "result": "pass", "after_photo_id": "f" * 32},
        )

        assert resp.status_code == 404
        assert hazards.fetch(hazard_no).status == hazards.STATUS_NOTIFIED

    def test_拿一份文书当复查照片_销不了项(self, client: TestClient) -> None:
        """🔴 §5.2 的红线是「拿不到**照片**就没有任何路径能把状态改成 closed」。

        只问「这个编号取得到 sidecar 吗」是不够的:我们自己签发的文书就是产物,
        它的 artifact_id 还大大方方摆在界面的下载卡上 —— 随手复制一个回填到复查那一栏,
        就能把隐患销项,而复查证据成了「我们自己出的那张纸」,现场一张照片都没有。
        """
        hazard_no = _arrive(hazards.STATUS_OPEN)
        签发 = client.post(
            "/supervision/notice", json={"hazard_no": hazard_no, "due_phrase": DUE_PHRASE}
        )
        assert 签发.status_code == 200
        文书编号 = 签发.json()["data"]["documents"][0]["artifact_id"]

        resp = client.post(
            "/supervision/reinspect-result",
            json={"hazard_no": hazard_no, "result": "pass", "after_photo_id": 文书编号},
        )

        assert resp.status_code == 400
        assert resp.json()["error_code"] == "INVALID_INPUT"
        assert "不是照片" in resp.json()["user_msg"]  # 说清是"拿错了东西",不是"编号找不到"
        assert "REPORT" not in resp.json()["user_msg"]  # 内部枚举值不许露头
        # 隐患没被销项,证据链里也没多出一条复查记录
        assert hazards.fetch(hazard_no).status == hazards.STATUS_NOTIFIED
        assert [r.doc_type for r in hazards.docs_of([hazard_no])] == ["notice"]

    def test_照片正文被清掉只剩元数据_也销不了项(self, client: TestClient) -> None:
        """产物是「正文 + sidecar」两个文件,清理器 / 手工删档 / 落盘半截都可能只剩后者。

        只读 sidecar 的话这种产物照样"存在",于是隐患拿一张**已经不存在的照片**销了项 ——
        等到要拿证据链去追责那天,复查那一格点开是空的。

        人话必须与「拿错编号」那条**不一样**:文件丢了该重新传一张,编号拿错了该去重找,
        给同一句话的话工友会一直核对一个本来就没错的编号。
        """
        hazard_no = _arrive(hazards.STATUS_NOTIFIED)
        photo_id = _photo()
        artifacts.resolve(photo_id).unlink()  # 只删正文,sidecar 留着

        resp = client.post(
            "/supervision/reinspect-result",
            json={"hazard_no": hazard_no, "result": "pass", "after_photo_id": photo_id},
        )

        assert resp.status_code == 404
        assert resp.json()["error_code"] == "NOT_FOUND"
        assert "重新传一张" in resp.json()["user_msg"]
        assert hazards.fetch(hazard_no).status == hazards.STATUS_NOTIFIED
        assert hazards.docs_of([hazard_no]) == []

    def test_结论不在词表被拒(self, client: TestClient) -> None:
        hazard_no = _arrive(hazards.STATUS_NOTIFIED)

        resp = client.post(
            "/supervision/reinspect-result",
            json={"hazard_no": hazard_no, "result": "大概行吧", "after_photo_id": _photo()},
        )

        assert resp.status_code == 400

    def test_没停过工的复查合格直接销项(self, client: TestClient) -> None:
        hazard_no = _arrive(hazards.STATUS_NOTIFIED)
        photo_id = _photo()

        resp = client.post(
            "/supervision/reinspect-result",
            json={"hazard_no": hazard_no, "result": "pass", "after_photo_id": photo_id},
        )

        assert resp.status_code == 200
        assert resp.json()["data"]["status"] == hazards.STATUS_CLOSED
        assert resp.json()["data"]["documents"] == []  # 复查不出文书
        rows = hazards.docs_of([hazard_no])
        assert [(r.doc_type, r.photo_id, r.result) for r in rows] == [
            ("reinspect", photo_id, "pass")
        ]

    def test_停过工的复查合格是待签复工令不是销项(self, client: TestClient) -> None:
        """Codex#6 点名的失效模式:停过工的从这里销项 = **漏发复工令**。
        挑边由 db 按 was_suspended 决定,端点一个字都不许自己挑。"""
        hazard_no = _arrive(hazards.STATUS_SUSPENDED, grade=hazards.GRADE_SEVERE, severity="重大")

        resp = client.post(
            "/supervision/reinspect-result",
            json={"hazard_no": hazard_no, "result": "pass", "after_photo_id": _photo()},
        )

        assert resp.status_code == 200
        assert resp.json()["data"]["status"] == hazards.STATUS_RESUMING
        assert "复工令" in resp.json()["user_msg"]
        assert hazards.fetch(hazard_no).closed_at is None  # 还没销

    def test_不合格进复查不合格并留照片(self, client: TestClient) -> None:
        hazard_no = _arrive(hazards.STATUS_NOTIFIED)

        resp = client.post(
            "/supervision/reinspect-result",
            json={"hazard_no": hazard_no, "result": "fail", "after_photo_id": _photo()},
        )

        assert resp.status_code == 200
        assert hazards.fetch(hazard_no).status == hazards.STATUS_REINSPECT_FAILED


# ---------------------------------------------------------------------------
# 复工令与上报:一条走完整闭环
# ---------------------------------------------------------------------------


class Test收尾两条路:
    def test_复工令签发之后销项且引用了原暂停令(self, client: TestClient) -> None:
        hazard_no = _arrive(hazards.STATUS_RESUMING, grade=hazards.GRADE_SEVERE, severity="重大")

        resp = client.post("/supervision/resume", json={"hazard_no": hazard_no})

        assert resp.status_code == 200
        data = resp.json()["data"]
        assert [d["doc_type"] for d in data["documents"]] == ["resumption"]
        assert data["documents"][0]["doc_no"].startswith("GYT-FG-")
        row = hazards.fetch(hazard_no)
        assert row.status == hazards.STATUS_CLOSED
        assert row.closed_at is not None

    def test_上报主管部门出监理报告并附证据链(self, client: TestClient) -> None:
        hazard_no = _arrive(hazards.STATUS_REINSPECT_FAILED)

        resp = client.post("/supervision/escalate", json={"hazard_no": hazard_no})

        assert resp.status_code == 200
        data = resp.json()["data"]
        assert data["status"] == hazards.STATUS_ESCALATED
        assert data["documents"][0]["doc_no"].startswith("GYT-JB-")
        assert hazards.fetch(hazard_no).status == hazards.STATUS_ESCALATED

    def test_全链走一遍_严重隐患从确认到复工(self, client: TestClient) -> None:
        """一条严重隐患的完整闭环:确认 → 三文书 → 复查合格 → 复工令 → 销项。

        逐段测过了还要有这一条,是因为**段与段之间的衔接**才是最容易断的地方
        (最典型的:复查合格之后到底该不该出复工令)。
        """
        hazard_no = _new_hazard(grade=hazards.GRADE_SEVERE, severity="重大", item="临边无防护")

        assert client.post("/supervision/confirm", json={"hazard_no": hazard_no}).status_code == 200
        assert (
            client.post(
                "/supervision/suspend",
                json={"hazard_no": hazard_no, "due_phrase": DUE_PHRASE},
            ).status_code
            == 200
        )
        assert (
            client.post(
                "/supervision/reinspect-result",
                json={"hazard_no": hazard_no, "result": "pass", "after_photo_id": _photo()},
            ).json()["data"]["status"]
            == hazards.STATUS_RESUMING
        )
        assert client.post("/supervision/resume", json={"hazard_no": hazard_no}).status_code == 200

        row = hazards.fetch(hazard_no)
        assert row.status == hazards.STATUS_CLOSED
        # 证据链:三份文书 + 一条复查 + 一份复工令,全在,顺序即发生先后
        assert [r.doc_type for r in hazards.docs_of([hazard_no])] == [
            "notice",
            "suspension",
            "owner_report",
            "reinspect",
            "resumption",
        ]
        assert len(_docx_files()) == 4  # 复查记录不出文书


# ---------------------------------------------------------------------------
# 请求体
# ---------------------------------------------------------------------------


def test_坏JSON回400而不是500(client: TestClient) -> None:
    """解析不了的请求体是用户侧的问题,不该被兜底成「系统开小差」。"""
    resp = client.post(
        "/supervision/notice",
        content=b"{not json at all",
        headers={"Content-Type": "application/json"},
    )

    assert resp.status_code == 400
    assert resp.json()["error_code"] == "INVALID_INPUT"


def test_信封永远是四键(client: TestClient) -> None:
    """成功与失败两条路都得是同一个形状 —— 前端只写一套归一化。"""
    hazard_no = _arrive(hazards.STATUS_OPEN)
    成功 = client.post(
        "/supervision/notice", json={"hazard_no": hazard_no, "due_phrase": DUE_PHRASE}
    ).json()
    失败 = client.post(
        "/supervision/notice", json={"hazard_no": "没这号", "due_phrase": "明天"}
    ).json()

    for body in (成功, 失败):
        assert set(body) == {"ok", "data", "user_msg", "error_code"}
    assert 成功["error_code"] is None
    assert 失败["data"] is None


# ===========================================================================
# W10:两条 GET + 否决 —— 造数、取数、直查库
# ===========================================================================

_清单行的键: Final[frozenset[str]] = frozenset(
    {
        # ① ``scoping.hazard_item()`` 的九个键(对话链念的也是这九个,同一份形状)
        "hazard_no",
        "item",
        "grade",
        "status",
        "status_display",
        "due_date",
        "due_display",
        "overdue",
        "needs_grading",
        # ② ``_row_payload`` 另加的两个:表格要、对话不要
        #    · severity —— **未定级的隐患对人念的是它**,不是 grade(那时 grade 是映射表给的
        #      默认档「一般」,不是有人判过的结论);
        #    · project_id —— 操作台会跨工地看,不给这一格就分不出哪条属于谁。
        "severity",
        "project_id",
        #    · photo_id —— 🔴 **发现这条隐患的那张现场照片**(2026-08-17 补)。
        #      补之前操作台上一条隐患只有文字,监理要在看不到照片的情况下判
        #      一般/严重、还要判是不是「识错了」而按下否决 —— 而否决这个判断
        #      完全依赖看照片。真人反馈只有三个字:「没有照片」。
        #      ⚠️ 不是复查照片(那张在 hazard_docs.photo_id,一次复查一张)。
        "photo_id",
    }
)
"""清单/详情里一行**恰好**有的键。

写成一份写死的集合(而不是从 ``hazard_item()`` 现算)是刻意的:现算的话,谁往
``hazard_item`` 里悄悄加一个键,这条断言会跟着一起变、什么都拦不住。而这个键集合是
**对外契约** —— 前端按它渲染,加字段要连前端一起改,减字段会让某一侧静默少一格。
下面那条用例另有一句反过来对着 ``hazard_item()`` 核一遍,两头都钉住。
"""

_证据链的键: Final[frozenset[str]] = frozenset(
    {
        "doc_type",
        "doc_type_display",
        "doc_no",
        "artifact_id",
        "filename",
        "photo_id",
        "result",
        "result_display",
        "created_at",
    }
)
"""``documents[*]`` 里一项**恰好**有的键。文书行与复查行**共用这一份形状**,差别只在
哪几格有值 —— 拆成两种形状的话前端得先猜自己拿到的是哪一种,而猜错的表现是少渲一块、
不报错。"""

_LIMIT_ENV: Final[str] = "GYT_SUPERVISION_LIST_MAX_ROWS"


def _调小上限(monkeypatch: pytest.MonkeyPatch, value: int) -> None:
    """把清单条数上限调小(写法与 ``test_supervision_tools._use_limit`` 同款)。

    ⚠️ 设完必须 ``get_settings.cache_clear()``:``get_settings`` 挂着 lru_cache,不清的话
    改的是环境变量、读的还是旧 Settings —— 表现是 ``truncated`` 恒 False,截断那条路
    一次都没被跑过而用例照绿。conftest 的 autouse fixture 会在用例前后各清一次,
    所以这里只清这一下就够,不用自己收拾环境。
    """
    monkeypatch.setenv(_LIMIT_ENV, str(value))
    get_settings.cache_clear()


def _库里还有(hazard_no: str) -> bool:
    """**直查 sqlite**:这一行到底还在不在。

    回执说什么不算,库里有没有才算 —— 本仓真机验收的口径(``live_acceptance.py``
    的"库级铁证"同一条)。刻意不走 ``hazards.fetch``:否决走的是 ``delete_pending``,
    与 ``fetch`` 同一个模块、同一套连接,拿它验等于让被告自己作证。
    """
    with closing(sqlite3.connect(get_settings().sqlite_path)) as conn:
        found = conn.execute("SELECT 1 FROM hazards WHERE hazard_no = ?", (hazard_no,)).fetchone()
    return found is not None


def _清单(client: TestClient, **params: str) -> dict[str, Any]:
    """打清单端点,断言 200,返回**整个信封**(``user_msg`` 也有用例要断,不能只给 data)。

    🔴 **没传的参数就是"没这个键"** —— ``project_id`` 三态里有两态靠这个分辨,
    所以这里绝不许给任何参数补默认值。要发「有键无值」那一态,走 ``_清单_原始查询串``。
    """
    resp = client.get("/supervision/hazards", params=params)
    assert resp.status_code == 200, resp.text
    return resp.json()


def _清单_原始查询串(client: TestClient, **params: str) -> dict[str, Any]:
    """同上,但查询串由 ``urlencode`` 自己拼。

    **「有键无值」(``?project_id=``)那一态只能走它**:那一态能不能表达出来,取决于
    HTTP 客户端怎么编码空值,而那不是我们该替它保证的事;``urlencode({"project_id": ""})``
    产出 ``project_id=`` 是标准库钉死的行为。
    """
    resp = client.get(f"/supervision/hazards?{urlencode(params)}")
    assert resp.status_code == 200, resp.text
    return resp.json()


def _详情(client: TestClient, hazard_no: str) -> dict[str, Any]:
    """打详情端点,断言 200,返回整个信封。"""
    resp = client.get(f"/supervision/hazards/{hazard_no}")
    assert resp.status_code == 200, resp.text
    return resp.json()


def _编号(信封: dict[str, Any]) -> set[str]:
    """这一页列出来的隐患编号集合。

    用集合是刻意的:这一组用例断的是"谁在谁不在",顺序另有 db 层的用例盯着
    (``hazards._ORDER_BY``),在这儿再断一遍只会让筛子的用例被排序的改动带红。
    """
    return {h["hazard_no"] for h in 信封["data"]["hazards"]}


class _台账(NamedTuple):
    """清单用例共用的那张台账。**每一条都有存在的理由,别顺手删:**

    · ``已销项`` / ``已上报`` —— 「全部」与「在办」的**唯一**差别就是这两档
      (``scoping.CLOSED_STATUSES``);没有它们,两个筛子返回同一批,
      「全部一条都不筛」与「缺省是在办」两条用例会双双变成恒绿。
    · ``今天到期`` —— 超期边界那条(严格小于)的全部依据。
    · ``未归属待确认`` / ``未归属超期`` —— ``project_id`` 三态靠它们才分得开;
      而且两条分属不同筛子,``unassigned`` 有没有被筛子筛过一眼看得出来。
    """

    待确认: str
    超期: str
    今天到期: str
    在办: str
    已销项: str
    已上报: str
    未归属待确认: str
    未归属超期: str

    @property
    def 有工地的(self) -> set[str]:
        """挂在 ``PROJECT`` 名下的那几条(= 全部减去未归属那两条)。"""
        return set(self) - self.未归属的

    @property
    def 未归属的(self) -> set[str]:
        return {self.未归属待确认, self.未归属超期}


@pytest.fixture
def 台账() -> _台账:
    """搭一张覆盖四个筛子 × 两种归属的台账。今天钉在 ``TODAY``(各用例的 autouse fixture 干的)。"""
    册 = _台账(
        待确认=_arrive(hazards.STATUS_PENDING, item="待确认的"),
        超期=_arrive(hazards.STATUS_NOTIFIED, due=YESTERDAY_ISO, item="昨天到期的"),
        今天到期=_arrive(hazards.STATUS_NOTIFIED, due=TODAY_ISO, item="今天到期的"),
        在办=_arrive(hazards.STATUS_OPEN, item="已确认待处置的"),
        已销项=_arrive(hazards.STATUS_CLOSED, item="已经销项的"),
        已上报=_arrive(hazards.STATUS_ESCALATED, item="已经上报的"),
        未归属待确认=_arrive(hazards.STATUS_PENDING, project_id=UNASSIGNED, item="没选工地时拍的"),
        未归属超期=_arrive(
            hazards.STATUS_NOTIFIED,
            due=YESTERDAY_ISO,
            project_id=UNASSIGNED,
            item="没选工地又超期的",
        ),
    )
    # 幂等键是 (project_id, photo_sha256, item),撞上了 create() 会静默回第一条那行 ——
    # 那样"我造了 8 条"其实只有几条,而下面的断言照样能凑合过去(见 _new_hazard 的警告)。
    assert len(set(册)) == len(册), "造数撞了幂等键,台账没造齐"
    return 册


@pytest.fixture
def _钉住今天(monkeypatch: pytest.MonkeyPatch) -> None:
    """把「今天」钉死在 ``TODAY``。

    ⚠️ **桩必须打在 ``scoping.today_hk`` 这个模块属性上。** ``scoping.py`` 那个函数的
    docstring 明写着:按名 ``from ... import today_hk`` 的调用点在导入那一刻就把函数对象
    绑死了,桩打不进去 —— 而**打不进去的表现不是红**,是超期那几条用例跟着跑测试的
    真实日子飘(今天绿、明天红,查起来极费劲)。
    所以下面有一条用例专门先断 ``data.today`` 等于钉的那天,再去断超期。
    """
    monkeypatch.setattr(scoping, "today_hk", lambda: TODAY)


# ---------------------------------------------------------------------------
# GET /supervision/hazards —— 清单
# ---------------------------------------------------------------------------


@pytest.mark.usefixtures("_钉住今天")
class Test隐患清单:
    def test_不传scope就是在办(self, client: TestClient, 台账: _台账) -> None:
        """缺省档选错的表现是:面板一打开就少一批(或多一批)隐患,而界面上看不出来。

        三边比:缺省 == 显式「在办」,且**两者都不等于「全部」**。少了最后那一句的话,
        缺省档哪天悄悄变成「全部」,前两条照样绿(台账里有已销项/已上报两条撑着差异)。
        """
        缺省 = _清单(client)
        显式 = _清单(client, scope=scoping.SCOPE_ACTIVE)
        全部 = _清单(client, scope=scoping.SCOPE_ALL)

        assert 缺省["data"]["scope"] == scoping.SCOPE_ACTIVE
        assert _编号(缺省) == _编号(显式)
        assert _编号(缺省) != _编号(全部), "缺省档已经不是「在办」了 —— 它现在筛的是别的东西"

    def test_全部一条都不筛_含已销项与已上报(self, client: TestClient, 台账: _台账) -> None:
        """「全部」就是全部。造数里**必须**有已销项与已上报两条,否则这条测了个寂寞:
        「全部」与「在办」的唯一差别就是这两档,没有它们两个筛子返回同一批。
        """
        全部 = _清单(client, scope=scoping.SCOPE_ALL)

        assert _编号(全部) == set(台账)
        assert 全部["data"]["total"] == len(台账)
        # 点名这两条:它们在「全部」里,不在「在办」里 —— 差异真的来自它们
        assert {台账.已销项, 台账.已上报} <= _编号(全部)
        assert {台账.已销项, 台账.已上报}.isdisjoint(_编号(_清单(client)))

    def test_待确认只出pending(self, client: TestClient, 台账: _台账) -> None:
        """D17 的人工确认闸就靠这一档把活推到人面前。混进别的状态 = 监理去确认一条
        早就签过文书的隐患,而端点会拒 —— 人只看得到一次莫名其妙的失败。"""
        待确认 = _清单(client, scope=scoping.SCOPE_PENDING)

        assert _编号(待确认) == {台账.待确认, 台账.未归属待确认}
        assert 待确认["data"]["total"] == 2
        assert [h["status"] for h in 待确认["data"]["hazards"]] == [hazards.STATUS_PENDING] * 2

    def test_超期这一档与scoping的判据等价(self, client: TestClient, 台账: _台账) -> None:
        """判据只有一份(``scoping.is_overdue``),端点不许自己再写一套 —— 抄第二份的表现是
        面板筛出来的条数与对话里念的对不上,而两边测试都绿(TODO-45 A 组那类漂移)。

        两头都断:既要与 ``is_overdue`` 逐条一致,**也要**等于人工数出来的那两条 ——
        只断前者的话,判据本身写错时两边一起错、还互相印证。
        """
        超期 = _清单(client, scope=scoping.SCOPE_OVERDUE)

        照判据算的 = {r.hazard_no for r in hazards.list_rows() if scoping.is_overdue(r, TODAY_ISO)}
        assert _编号(超期) == 照判据算的
        assert 照判据算的 == {台账.超期, 台账.未归属超期}

    def test_今天到期的不算超期(self, client: TestClient, 台账: _台账) -> None:
        """🔴 边界是**严格小于**。差一天的后果不是"少列一条",是把还在期限内的施工方
        写成「拒不整改」—— 而那句话要进《监理报告》报建设主管部门。

        两头都断:它不在超期清单里,**而且**它自己那一行的 ``overdue`` 是 False。
        只断前者的话,把 ``<`` 改成 ``<=`` 时清单红了、行里那一格却没人管。
        再加一条"昨天到期的确实算超期",免得整条用例在超期功能整个坏掉时也恒绿。
        """
        超期 = _清单(client, scope=scoping.SCOPE_OVERDUE)
        assert 台账.今天到期 not in _编号(超期)

        行 = {h["hazard_no"]: h for h in _清单(client, scope=scoping.SCOPE_ALL)["data"]["hazards"]}
        assert 行[台账.今天到期]["due_date"] == TODAY_ISO, "造数造歪了,它不是今天到期"
        assert 行[台账.今天到期]["overdue"] is False
        assert 行[台账.超期]["overdue"] is True

    def test_野筛子被拒且回的话里一个英文字母都没有(self, client: TestClient) -> None:
        """``scoping.in_scope`` 认不出的词会**静默落在「在办」**且一声不吭(它的 docstring
        点名要求调用方自己拦野词),所以拦是端点的责任 —— 不拦的话少给的清单在界面上
        看不出来少了。

        顺带盯住这句话是**说给工地上的人听的**:不许漏出 scope / SCOPES / KeyError
        这类内部词(写法照 ``test_期限看不懂就拒`` 那条)。
        """
        resp = client.get("/supervision/hazards", params={"scope": "严重的"})

        assert resp.status_code == 400
        body = resp.json()
        assert body["error_code"] == "INVALID_INPUT"
        assert "严重的" in body["user_msg"], "得告诉人他说的是哪个词"
        for 词 in scoping.SCOPES:
            assert 词 in body["user_msg"], "也得告诉人有哪几个词可以用"
        assert not re.search(r"[A-Za-z]", body["user_msg"]), "回给工地的话里冒出了英文"

    def test_project_id不出现就是全部工地(self, client: TestClient, 台账: _台账) -> None:
        """三态之一。``list_rows(project_id=None)`` = 不筛工地。"""
        data = _清单(client, scope=scoping.SCOPE_ALL)["data"]

        assert data["project_id"] is None
        assert _编号({"data": data}) == set(台账)

    def test_project_id有键无值只看未归属(self, client: TestClient, 台账: _台账) -> None:
        """三态之二。``?project_id=`` → ``list_rows(project_id="")`` = 只看未归属(D6)。"""
        信封 = _清单_原始查询串(client, scope=scoping.SCOPE_ALL, project_id=UNASSIGNED)

        assert 信封["data"]["project_id"] == UNASSIGNED
        assert _编号(信封) == 台账.未归属的
        assert all(h["project_id"] == UNASSIGNED for h in 信封["data"]["hazards"])

    def test_project_id给了工地就只看那个工地(self, client: TestClient, 台账: _台账) -> None:
        """三态之三。顺带断一句 ``user_msg``:未归属那批**不在**这份清单里,
        不专门说一句就永远没人看见(D6 —— 那批本来就是最容易没人管的)。"""
        信封 = _清单(client, scope=scoping.SCOPE_ALL, project_id=PROJECT)

        assert 信封["data"]["project_id"] == PROJECT
        assert _编号(信封) == 台账.有工地的
        assert all(h["project_id"] == PROJECT for h in 信封["data"]["hazards"])
        assert "另外还有 2 条隐患没归到任何工地" in 信封["user_msg"]

    def test_不出现与有键无值必须是两种东西(self, client: TestClient, 台账: _台账) -> None:
        """🔴 **这一组里最要紧的一条。**

        两态被写成同一个(``params.get(key, "")``、或者后面接一个 ``or ""``)的后果:
        监理打开操作台只看得见未归属那一堆,**而界面上一切正常** —— 没有报错、没有空列表、
        每一条都长得对,只是有工地的那一大半凭空不见了。

        所以断的不是"某个字段等于什么",而是**两次请求的结果必须不同**,并且差的正好是
        「有工地的那一批」。台账里同时有归属与未归属两种,两态才天然不同。
        """
        不出现 = _清单(client, scope=scoping.SCOPE_ALL)
        有键无值 = _清单_原始查询串(client, scope=scoping.SCOPE_ALL, project_id=UNASSIGNED)

        assert 不出现["data"]["project_id"] is None
        assert 有键无值["data"]["project_id"] == UNASSIGNED
        assert 不出现["data"]["total"] != 有键无值["data"]["total"]
        assert _编号(有键无值) < _编号(不出现), "「未归属」应当是「全部工地」的真子集"
        assert _编号(不出现) - _编号(有键无值) == 台账.有工地的, (
            "两态塌成一个了 —— 「全部工地」现在只给未归属那一堆"
        )

    def test_三态数未归属的算法不同但答的是同一个数(self, client: TestClient, 台账: _台账) -> None:
        """``unassigned`` 回答的是「有没有一批隐患没人看得见」(D6),三态各有各的算法:

            不出现   → 在已经取回的 rows 里数(不打第二次库)
            有键无值 → rows 本身就是未归属那一堆,数长度
            具体工地 → 它们**不在** rows 里,只能另查一次(只数条数)

        三条路必须答出同一个数。最容易漏的是第三支:漏了的话「另外还有 N 条没归到任何
        工地」这句话从此不出现,而 D6 那批就此没人看得见 —— 一句报错都不会有。
        """
        assert _清单(client, scope=scoping.SCOPE_ALL)["data"]["unassigned"] == 2
        无值 = _清单_原始查询串(client, scope=scoping.SCOPE_ALL, project_id=UNASSIGNED)
        assert 无值["data"]["unassigned"] == 2
        assert _清单(client, scope=scoping.SCOPE_ALL, project_id=PROJECT)["data"]["unassigned"] == 2

    def test_未归属条数不过scope筛子(self, client: TestClient, 台账: _台账) -> None:
        """``unassigned`` **不过筛子**:拿筛子筛过反而会把它藏起来。

        最容易看出来的是「超期」这一档:超期清单里只有 1 条未归属,而未归属一共 2 条。
        过了筛子的话这个数会变成 1,而屏幕上只是一个小数字变了 —— 没人会发现。
        """
        for scope in scoping.SCOPES:
            assert _清单(client, scope=scope)["data"]["unassigned"] == 2, (
                f"scope={scope} 时未归属条数被筛子改了"
            )

        超期 = _清单(client, scope=scoping.SCOPE_OVERDUE)
        assert sum(1 for h in 超期["data"]["hazards"] if not h["project_id"]) == 1
        assert 超期["data"]["unassigned"] == 2

    def test_计数算在筛出来的全部行上而不是截断后那批(
        self, client: TestClient, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """监理要的是「一共还有几条」,截断只影响列出来几行。

        算在截断后那批上的表现:台账里超期 30 条,面板上说"超期 2 条" —— 而那 28 条
        没人会去催。``truncated`` 同样不许省:不如实说的话,监理会以为台账里就这么几条。
        """
        limit = 2
        _调小上限(monkeypatch, limit)
        超期们 = [
            _arrive(hazards.STATUS_NOTIFIED, due=YESTERDAY_ISO, item=f"超期{i}") for i in range(3)
        ]
        待确认 = _arrive(hazards.STATUS_PENDING, item="等着确认的")
        assert len({*超期们, 待确认}) == 4, "造数撞了幂等键,没造出 4 条"

        data = _清单(client, scope=scoping.SCOPE_ALL)["data"]

        assert data["truncated"] is True
        assert len(data["hazards"]) == limit
        assert data["total"] == 4
        assert data["overdue"] == 3
        assert data["pending"] == 1

    def test_一行恰好这十二个键(self, client: TestClient, 台账: _台账) -> None:
        """键集合是**对外契约**:前端按它渲染,加字段要连前端一起改,减字段会让某一侧
        静默少一格(比如没了 ``severity``,未定级的隐患在面板上会被念成「一般」——
        而硬拦③ ``_require_graded`` 拦的正是这句话)。

        两头都钉:先对着写死的那份集合,再反过来核一句"其中九个真的来自
        ``scoping.hazard_item()``" —— 只写前者,``hazard_item`` 那边悄悄挪走一个键、
        端点这边补一个同名的,契约看着没变而两条链已经不同源了。
        """
        data = _清单(client, scope=scoping.SCOPE_ALL)["data"]
        assert data["hazards"], "台账是空的,这条用例什么都没验到"
        for 行 in data["hazards"]:
            assert set(行) == _清单行的键

        源头 = hazards.fetch(台账.在办)
        assert set(scoping.hazard_item(源头, today_iso=TODAY_ISO)) == _清单行的键 - {
            "severity",
            "project_id",
            "photo_id",
        }

    def test_今天是钉住的那天且超期真按它算(
        self, client: TestClient, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """``overdue`` 必须是**真算出来的**,不是造数时写死的一格。

        做法:同一条隐患、同一份库,只把「今天」拨来拨去,看它跟不跟着变。

        ⚠️ 先断 ``data.today`` 等于钉的那天 —— **这一句是在验桩本身生效了**。桩要是
        没打进去(比如哪天有人把端点改成 ``from ... import today_hk``),下面的超期断言
        会跟着跑测试的真实日子飘:今天绿、明天红,而没人会想到是桩没打进去。
        ``today`` 这个字段本来就是端点原样给出去的("为什么这条标了超期"要能解释),
        正好当桩的见证。
        """
        hazard_no = _arrive(hazards.STATUS_NOTIFIED, due=YESTERDAY_ISO, item="八月十九到期")

        monkeypatch.setattr(scoping, "today_hk", lambda: date(2026, 8, 20))
        晚 = _清单(client, scope=scoping.SCOPE_ALL)["data"]
        assert 晚["today"] == "2026-08-20", "桩没生效 —— 下面的断言全在跟真实日子赌运气"
        assert 晚["hazards"][0]["overdue"] is True
        assert 晚["overdue"] == 1
        assert _编号(_清单(client, scope=scoping.SCOPE_OVERDUE)) == {hazard_no}

        # 同一条隐患,把「今天」拨到期限之前 —— 它就不该再是超期
        monkeypatch.setattr(scoping, "today_hk", lambda: date(2026, 8, 18))
        早 = _清单(client, scope=scoping.SCOPE_ALL)["data"]
        assert 早["today"] == "2026-08-18"
        assert 早["hazards"][0]["overdue"] is False
        assert 早["overdue"] == 0
        assert _编号(_清单(client, scope=scoping.SCOPE_OVERDUE)) == set()

    def test_空台账是200加人话_不是错误(self, client: TestClient) -> None:
        """查空不是失败(口径对齐 schedule 的「台账里现在没有任务」)。"""
        信封 = _清单(client)

        assert 信封["ok"] is True
        assert 信封["data"]["total"] == 0
        assert 信封["data"]["hazards"] == []
        assert "一条都没有" in 信封["user_msg"]


# ---------------------------------------------------------------------------
# GET /supervision/hazards/{hazard_no} —— 详情 + 证据链
# ---------------------------------------------------------------------------


def _复查(client: TestClient, hazard_no: str, result: str, photo_id: str) -> None:
    """登记一次复查结论(走真端点,不绕 db)。"""
    resp = client.post(
        "/supervision/reinspect-result",
        json={"hazard_no": hazard_no, "result": result, "after_photo_id": photo_id},
    )
    assert resp.status_code == 200, resp.text


@pytest.mark.usefixtures("_钉住今天")
class Test隐患详情:
    def _签发一份通知单(self, client: TestClient) -> tuple[str, dict[str, Any]]:
        """搭一条已经签过《监理通知单》的隐患,返回(编号, 签发回执的 data)。

        **走真端点签发**,不用 db 直推:这一组里最要紧的一条(文件名同源)要拿签发回执
        当比对基准,db 直推出来的那份 ``doc_no`` / ``artifact_id`` 是测试自己编的,
        比对了也不说明两条真实代码路径同源。
        """
        hazard_no = _arrive(hazards.STATUS_OPEN, item="临边无防护")
        resp = client.post(
            "/supervision/notice", json={"hazard_no": hazard_no, "due_phrase": DUE_PHRASE}
        )
        assert resp.status_code == 200, resp.text
        return hazard_no, resp.json()["data"]

    def test_详情字段齐全且清单那一行原样在里面(self, client: TestClient) -> None:
        """详情 = 清单那一行 + 四个键。少一个前端就少渲一块,而且不会报错。"""
        hazard_no, _ = self._签发一份通知单(client)

        信封 = _详情(client, hazard_no)
        data = 信封["data"]

        assert set(data) == set(_清单行的键) | {
            "found_at",
            "closed_at",
            "reinspected",
            "documents",
        }
        assert data["hazard_no"] == hazard_no
        assert data["status"] == hazards.STATUS_NOTIFIED
        assert data["status_display"] == "已签发通知单"
        assert data["project_id"] == PROJECT
        assert data["found_at"]
        assert data["closed_at"] is None  # 还没销项
        assert "已签 1 份文书" in 信封["user_msg"]

    def test_文书行与复查行是两种形状(self, client: TestClient) -> None:
        """两种形状共用一个函数,差别只在哪几格有值。

        **这条必须同时拿到两种行**:只造文书的话,``photo_id`` / ``result`` 那几格恒为
        None,把它们写反(比如复查行也去填 filename)照样绿。
        """
        hazard_no, 回执 = self._签发一份通知单(client)
        复查照片 = _photo("整改后.jpg")
        _复查(client, hazard_no, "fail", 复查照片)

        documents = _详情(client, hazard_no)["data"]["documents"]

        assert [d["doc_type"] for d in documents] == ["notice", "reinspect"]
        for 项 in documents:
            assert set(项) == _证据链的键
            assert 项["created_at"]
        文书行, 复查行 = documents
        # 文书行:有件可下,没有复查那两格
        assert 文书行["artifact_id"] == 回执["documents"][0]["artifact_id"]
        assert 文书行["doc_type_display"] == "监理通知单"
        assert 文书行["filename"].endswith(".docx")
        assert 文书行["photo_id"] is None
        assert (文书行["result"], 文书行["result_display"]) == (None, None)
        # 复查行:反过来 —— 没有文件(它进不了下载卡),有照片有结论
        assert 复查行["artifact_id"] is None
        assert 复查行["filename"] is None, "复查记录没有文件,给了名字就是一张点开 404 的卡"
        assert 复查行["doc_type_display"] == "复查记录"
        assert 复查行["photo_id"] == 复查照片
        assert 复查行["result"] == "fail"

    def test_查不到编号是404且说的是查询那条路的话(self, client: TestClient) -> None:
        """措辞与 ``tools.get_hazard`` 那句一字不差,与写入那条路(``_require_hazard``)
        刻意不同:面板上点开一条隐患和在对话里问同一条是**同一个动作**,两处说法不一样
        会让人以为查的是两个台账;而签发时找不到是另一回事,那句在动作的语境里说。
        """
        resp = client.get("/supervision/hazards/GYT-H-20260101-000000-ffff")

        assert resp.status_code == 404
        body = resp.json()
        assert body["error_code"] == "NOT_FOUND"
        assert "台账里没有" in body["user_msg"]
        assert "没找到隐患" not in body["user_msg"], "这是写入那条路的话,别顺手统一了"

    def test_详情里的文件名与签发回执逐字相同(self, client: TestClient) -> None:
        """🔴 两处都走 ``_filename()``,谁也不许现拼第二份。

        不同源的表现是界面上下载卡的文件名和实际落盘的对不上,**而不会有任何报错**:
        卡片照渲、文件照下,只是名字换了一个 —— 等有人拿着文件名去盘上对账那天才发现。

        ⚠️ 断言必须是**逐字比对**。写成 ``endswith(".docx")`` 的话两边各拼各的也能过,
        那种断言在这里等于没写。下面第二句再钉一次"名字里带着编号那一段",
        免得两边同时退化成一个常量字符串也能过第一句。
        """
        hazard_no, 回执 = self._签发一份通知单(client)

        签发时的 = 回执["documents"][0]["filename"]
        详情里的 = _详情(client, hazard_no)["data"]["documents"][0]["filename"]

        assert 详情里的 == 签发时的, "详情与签发两条路的文件名不同源了"
        assert 回执["documents"][0]["doc_no"] in 详情里的
        assert 详情里的.endswith(".docx")

    def test_复查行的照片是复查那张不是发现那张(self, client: TestClient) -> None:
        """🔴 ``documents[*].photo_id`` 是 ``hazard_docs.photo_id``(**这一次复查**拍的),
        不是 ``hazards.photo_id``(首次发现那张)。

        混起来 = 拿发现时的照片当"整改后"的证据,而「复查必须挂照片」这条红线
        (方案 §5.2)的全部意义,就是事后追责时分得清这两张。

        造数时两张必须是**两个不同的 artifact_id** —— 同一张的话取错那一张也照样相等,
        这条用例就白写了(所以下面先断一句它们不同)。
        """
        hazard_no, _ = self._签发一份通知单(client)
        发现时那张 = hazards.fetch(hazard_no).photo_id
        复查那张 = _photo("整改后.jpg")
        assert 复查那张 != 发现时那张, "两张照片是同一个编号,这条用例验不出取错"

        _复查(client, hazard_no, "fail", 复查那张)
        documents = _详情(client, hazard_no)["data"]["documents"]

        复查行 = [d for d in documents if d["doc_type"] == "reinspect"]
        assert [d["photo_id"] for d in 复查行] == [复查那张]
        assert 发现时那张 not in {d["photo_id"] for d in documents}, (
            "拿发现时的照片当整改后的证据了"
        )

    def test_复查结论的中文名两个方向都对而文书行没有(self, client: TestClient) -> None:
        """留档文书上、回给工友的话里都不许出现 ``pass`` / ``fail`` 这种英文枚举值。

        三种格子一次断完:合格、不合格、**文书行是 None**(不是「—」——「—」是留档
        文书表格里"这格本来就空"的写法,JSON 里 null 才是诚实的"没有值",
        空格渲成什么由前端决定)。

        造两次复查(先不合格再合格)是必须的:一次只验得了其中一个方向。
        """
        hazard_no, _ = self._签发一份通知单(client)
        _复查(client, hazard_no, "fail", _photo("第一次复查.jpg"))
        _复查(client, hazard_no, "pass", _photo("第二次复查.jpg"))

        documents = _详情(client, hazard_no)["data"]["documents"]

        assert [(d["doc_type"], d["result"], d["result_display"]) for d in documents] == [
            ("notice", None, None),
            ("reinspect", "fail", "不合格"),
            ("reinspect", "pass", "合格"),
        ]

    def test_复查过没有这一格两种情况都对(self, client: TestClient) -> None:
        """「这条复查了吗」是监理翻这一页最常问的一句,只把它埋进 ``documents`` 里
        让人自己数,等于没答。

        ⚠️ **别拿 status 推**:复查不合格之后还能再复查,而 closed 也可能是复工令签出来的
        —— 状态答不了这个问题。所以这条两种情况都得跑一遍。
        """
        hazard_no, _ = self._签发一份通知单(client)

        没复查过 = _详情(client, hazard_no)
        assert 没复查过["data"]["reinspected"] is False
        assert "还没复查过" in 没复查过["user_msg"]

        _复查(client, hazard_no, "fail", _photo("整改后.jpg"))

        复查过 = _详情(client, hazard_no)
        assert 复查过["data"]["reinspected"] is True
        assert "复查过 1 次,最近一次不合格" in 复查过["user_msg"]


# ---------------------------------------------------------------------------
# POST /supervision/reject —— 否决一条待确认的隐患
# ---------------------------------------------------------------------------


class Test否决:
    def test_否决成功之后库里那一行真的没了(self, client: TestClient) -> None:
        """🔴 **回执说什么不算,库里有没有才算**(本仓真机验收的口径)。

        ``reject`` 是 ``db.delete_pending`` 唯一的调用入口(TODO-45 B 组:W9 落地时它
        零调用点,界面上的「否决」只是把那一行本地划掉,刷新就回来了)—— 这条要防的
        正是那件事重演:回执一切正常,而库里那行原封不动。

        ``data`` 用整体相等断:三个键一个不少(冻结形状),``status`` 是那个
        **不属于 db.STATUSES 八档**的哨兵值 —— 那一行已经删了,压根没有状态可报。
        """
        hazard_no = _arrive(hazards.STATUS_PENDING, item="拍糊了的那张")
        assert _库里还有(hazard_no), "造数没落库,下面那条断言会恒绿"

        resp = client.post("/supervision/reject", json={"hazard_no": hazard_no})

        assert resp.status_code == 200
        assert resp.json()["data"] == {
            "hazard_no": hazard_no,
            "status": supervision_api._STATUS_DELETED,
            "documents": [],
        }
        assert not _库里还有(hazard_no)
        assert hazards.fetch(hazard_no) is None

    def test_缺隐患编号是400(self, client: TestClient) -> None:
        """空 body 不许被当成"否决点了个啥都没有"放过去。"""
        resp = client.post("/supervision/reject", json={})

        assert resp.status_code == 400
        assert resp.json()["error_code"] == "INVALID_INPUT"
        assert "隐患编号" in resp.json()["user_msg"]

    def test_编号查不到是404(self, client: TestClient) -> None:
        resp = client.post("/supervision/reject", json={"hazard_no": "GYT-H-20260101-000000-ffff"})

        assert resp.status_code == 404
        assert resp.json()["error_code"] == "NOT_FOUND"

    def test_确认过的否决不了而且那一行还在(self, client: TestClient) -> None:
        """🔴 「拒了」和「拒了但也把行删了」在回执上长得一模一样,**只有直查库分得开**。

        留档与证据链不能因为一次误点消失:``delete_pending`` 的 ``WHERE status='pending'``
        是硬守卫,端点这道先读只为说人话(它能说清「这条现在是已确认待处置」)。
        """
        hazard_no = _arrive(hazards.STATUS_OPEN, item="已经确认过的")

        resp = client.post("/supervision/reject", json={"hazard_no": hazard_no})

        assert resp.status_code == 409
        assert resp.json()["error_code"] == "CONFLICT"
        assert "已确认待处置" in resp.json()["user_msg"], "得说清它现在在哪一站"
        assert _库里还有(hazard_no)
        assert hazards.fetch(hazard_no).status == hazards.STATUS_OPEN

    def test_先读之后被别人确认掉_否决必须落空(
        self, client: TestClient, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """并发窗口:端点是「先读 status 判断能不能否 → 再 DELETE」两步。

        这里不真起两个线程,手工在两步之间把库改掉 —— 效果一样而且必然复现:读到
        ``pending``(可以否决)之后、DELETE 之前,另一个人把它确认了。这时若照删,
        一条已经进入正式流程的隐患凭空消失,**而且一声不吭**。
        真正说了算的是 ``delete_pending`` 的返回值,先读的那份快照说了不算。
        """
        hazard_no = _arrive(hazards.STATUS_PENDING, item="正被人同时处理的")
        真的 = hazards.delete_pending

        def 抢在前面(no: str) -> bool:
            hazards.confirm(no)  # 「另一个请求」:先读之后、这条 DELETE 之前把它确认了
            return 真的(no)

        monkeypatch.setattr(hazards, "delete_pending", 抢在前面)

        resp = client.post("/supervision/reject", json={"hazard_no": hazard_no})

        assert resp.status_code == 409
        assert "刚被改过" in resp.json()["user_msg"]  # 人话:让人刷新再看
        assert _库里还有(hazard_no)
        assert hazards.fetch(hazard_no).status == hazards.STATUS_OPEN

    def test_挂了留档材料的删不动_收敛成人话而不是500(
        self, client: TestClient, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """``hazard_docs`` 有外键指着 ``hazards``。pending 的行理论上不该挂着任何文书
        (所有签发都要求 open 起步),但真撞上了得收敛成人话:让它变成 500 的话,
        工友看到的是「系统开小差」,而这其实是一句明确的「这条已经有留档了,删不得」。

        异常原文只进日志 —— ``FOREIGN KEY constraint failed`` 不是工地上的话。
        """
        hazard_no = _arrive(hazards.STATUS_PENDING, item="挂着留档的")

        def 撞外键(_no: str) -> bool:
            raise sqlite3.IntegrityError("FOREIGN KEY constraint failed")

        monkeypatch.setattr(hazards, "delete_pending", 撞外键)

        resp = client.post("/supervision/reject", json={"hazard_no": hazard_no})

        assert resp.status_code == 409
        assert resp.json()["error_code"] == "CONFLICT"
        assert "留档材料" in resp.json()["user_msg"]
        assert "FOREIGN KEY" not in resp.json()["user_msg"], "内部细节不许露头"
        assert _库里还有(hazard_no)

    def test_deleted这个哨兵值不许撞进db的八档(self) -> None:
        """``"deleted"`` 是全项目唯一一个**不属于** ``db.STATUSES`` 的状态值 —— 那一行已经
        从库里删掉了,压根没有状态可报,而回执形状是冻结的(``status`` 这个键必须在)。

        撞上了的表现:前端分不清「这条被否决了」和「这条处于 deleted 状态」,
        而两种情况该做的事完全相反(刷新划掉 vs 继续处置)。模块里有一道导入期硬失败
        守着,这条是把那个约束写成人看得见的形式。

        ⚠️ 拿常量断,**别手写 ``"deleted"``** —— 手写的那份在有人换了哨兵值时会照绿。
        """
        assert supervision_api._STATUS_DELETED not in hazards.STATUSES


# ---------------------------------------------------------------------------
# POST /supervision/photo —— 直传复查照片,换一个 photo_id
# ---------------------------------------------------------------------------
#
# 这条端点存在的理由:「登记复查结论」要填一串 32 位十六进制,而后端此前**没有任何
# 通用的传图口子** —— PHOTO 产物只能经由聊天的 core/uploads.py 产生。于是做复查的人
# 手机里刚拍完那张照片,却得先发进聊天框、再把图底下那串 hex 抄回表单。
# 真人测试的反馈原话是「照片不能是编号意义不明」。
#
# 这一组要钉死的静默错误:
#   · **只信 Content-Type** —— 这张照片接下来会被 ``_require_photo`` 当成复查证据,
#     而复查合格是**销项**的唯一通道。只信自述的话,一个 .txt 改个头就能把隐患销掉,
#     而留档文书上写着「隐患已消除」。魔数那道闸没了,现有用例**一条都不会红**
#     (它们传的本来就是合法图),所以这一组必须自己造反例。
#   · **只认 RIFF 就当 WebP** —— WAV / AVI 同为 RIFF 容器,会被当照片收下。
#   · **被拒的请求也落盘** —— 400/413/401 在回执上和"拒了但也存了"长得一模一样。
#   · **两个端点各自绿却接不上** —— 上传回的 photo_id 若过不了 ``_require_photo``
#     的三道校验(kind=PHOTO、sidecar 在、正文文件在),这条链就是断的,
#     而单独验任一端都看不出来。下面 ``test_换回来的编号复查端点真的认`` 串起来验。


def _传照片(client: TestClient, payload: bytes, **kwargs: Any) -> Any:
    """打上传端点。``content=`` 发的是**原始字节**(不是 multipart、不是 JSON)。"""
    return client.post("/supervision/photo", content=payload, **kwargs)


class Test照片直传:
    def test_传一张JPEG换回一个照片编号(self, client: TestClient) -> None:
        """快乐路径,把回执形状整个钉住。

        ``filename`` 要是**人话**(不是内部路径、不是 32 位 hex):它会原样进产物 sidecar,
        将来下载卡上显示的就是它。扩展名由**魔数**决定,不是客户端说了算。
        """
        resp = _传照片(client, FAKE_JPEG)

        assert resp.status_code == 200
        body = resp.json()
        assert body["ok"] is True
        assert body["error_code"] is None
        assert body["user_msg"], "得给工友一句话,不能空着"
        photo_id = body["data"]["photo_id"]
        assert re.fullmatch(r"[0-9a-f]{32}", photo_id), f"照片编号不是 32 位小写 hex:{photo_id!r}"
        assert body["data"]["filename"] == "复查照片.jpg"
        # 落盘的是 PHOTO 产物,而且正文字节一个不差 —— register 收的是原始 body。
        assert artifacts.read_meta(photo_id)["kind"] == ArtifactKind.PHOTO.value
        assert artifacts.resolve(photo_id).read_bytes() == FAKE_JPEG
        assert artifacts.resolve(photo_id).suffix == ".jpg", "没有后缀的照片浏览器打不开"

    def test_换回来的编号复查端点真的认(self, client: TestClient) -> None:
        """🔴 **这条是这一组里最强的一条** —— 它把上传与复查两个端点串起来验。

        单独验任一个都可能各自绿而接不上:``_require_photo`` 有三道校验
        (kind 必须是 PHOTO、sidecar 取得到、正文文件真的在),上传那边只要漏配一样
        (比如 kind 写成 DOCUMENT、或者扩展名被剥成空串导致正文落在另一个名字上),
        表现就是**上传成功、复查却说"没找到这张复查照片"** —— 而两边的用例各自都绿。

        走完整条真实动线:传照片 → 拿编号去登记复查合格 → 隐患销项,
        并且证据链里那条 reinspect 记录挂的正是这个编号。
        """
        hazard_no = _arrive(hazards.STATUS_NOTIFIED)

        上传 = _传照片(client, FAKE_JPEG)
        assert 上传.status_code == 200
        photo_id = 上传.json()["data"]["photo_id"]

        复查 = client.post(
            "/supervision/reinspect-result",
            json={"hazard_no": hazard_no, "result": "pass", "after_photo_id": photo_id},
        )

        assert 复查.status_code == 200, f"上传换来的编号复查端点不认:{复查.json()}"
        assert 复查.json()["data"]["status"] == hazards.STATUS_CLOSED
        assert [(r.doc_type, r.photo_id) for r in hazards.docs_of([hazard_no])] == [
            ("reinspect", photo_id)
        ]

    @pytest.mark.parametrize(
        ("样本", "扩展名"),
        [(FAKE_JPEG, ".jpg"), (FAKE_PNG, ".png"), (FAKE_WEBP, ".webp")],
        ids=["jpeg", "png", "webp"],
    )
    def test_三种格式都认而且落盘扩展名跟着魔数走(
        self, client: TestClient, 样本: bytes, 扩展名: str
    ) -> None:
        """JPG / PNG / WebP 三种。扩展名必须落在 ``config.ALLOWED_IMAGE_EXT`` 里,
        否则 ``artifacts._safe_ext`` 会把它剥成空串 —— 文件落盘没有后缀、浏览器打不开,
        而 ``_require_photo`` 只查 kind 和文件在不在,**照样放行**。
        """
        resp = _传照片(client, 样本)

        assert resp.status_code == 200, resp.json()
        photo_id = resp.json()["data"]["photo_id"]
        assert resp.json()["data"]["filename"].endswith(扩展名)
        assert artifacts.resolve(photo_id).suffix == 扩展名

    @pytest.mark.parametrize(
        ("样本", "叫什么"),
        [(FAKE_DOCX, "docx"), (FAKE_TEXT, "纯文本"), (FAKE_WAV, "WAV(同为RIFF容器)")],
    )
    def test_不是图片的一律拒而且盘上不留东西(
        self, client: TestClient, 样本: bytes, 叫什么: str
    ) -> None:
        """🔴 复查合格是**销项**的唯一通道 —— 放一个非图片文件进去,等于让人拿一个 .txt 销项。

        WAV 那一档单独存在:它和 WebP 一样以 ``RIFF`` 开头,只比前四个字节的实现会把它
        当照片收下(``_WEBP_TAG_AT`` 那段注释说的就是这件事)。

        「盘上不留东西」是另一半:被拒的请求先落盘再报 400 的话,产物目录会被随手一传
        就撑大,而回执看起来完全正常。这也钉住了实现里"先判类型、再落盘"的顺序。
        """
        之前 = _产物文件()

        resp = _传照片(client, 样本)

        assert resp.status_code == 400, f"{叫什么}被当成照片收下了:{resp.json()}"
        assert resp.json()["error_code"] == "INVALID_INPUT"
        assert "照片" in resp.json()["user_msg"]
        assert _产物文件() == 之前, f"{叫什么}被拒了,但盘上多了东西"

    def test_Content_Type说是图片也不算数(self, client: TestClient) -> None:
        """🔴 **判据是魔数,不是客户端自述的 Content-Type。**

        这条与上面那组的区别:上面传的反例不带任何声明,而**真正的攻击面是带着
        ``image/jpeg`` 头的非图片** —— 只信 header 的实现在上面那组里会红一部分
        (httpx 默认不给 content= 加 Content-Type),但在这一条上必红。
        「浏览器给的 MIME 不可靠」本仓在 DXF 那条线上已经证过一次。

        另断一句:回执里**不许复述**客户端报的那个 MIME —— 念出来只会让人以为
        "我明明填的是 image/jpeg 啊",而那个值正是我们不信的东西(它走 detail 进日志)。
        """
        之前 = _产物文件()

        resp = _传照片(client, FAKE_DOCX, headers={"Content-Type": "image/jpeg"})

        assert resp.status_code == 400
        assert "image/jpeg" not in resp.json()["user_msg"]
        assert _产物文件() == 之前

    def test_超过上限的照片回413而且盘上不留东西(
        self, client: TestClient, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """上限走 ``settings.photo_max_mb``(**与打卡同一个旋钮**,禁止硬编码)——
        「环境变量能改到它」本身就是断言的一半:改不到就说明有人把上限写死了。

        把上限压到 ~100 字节来测,免得真造 10MB(判据抄 ``test_checkin_api``
        的 ``test_超限body当场413且一步不往下走``)。
        """
        monkeypatch.setenv("GYT_PHOTO_MAX_MB", "0.0001")  # ≈104 字节
        get_settings.cache_clear()
        之前 = _产物文件()

        resp = _传照片(client, FAKE_JPEG + b"x" * 4096)

        assert resp.status_code == 413
        assert resp.json()["error_code"] == "FILE_TOO_LARGE"
        assert "10" not in resp.json()["user_msg"], "兆数得跟着 settings 走,不是写死的 10"
        assert _产物文件() == 之前, "超限的照片不许落盘"

    def test_没超上限的照片照收(self, client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
        """上一条的反面。没有它的话,「上限调成 0 字节」这种改法也能让上一条绿,
        而那等于所有照片都传不上来。"""
        monkeypatch.setenv("GYT_PHOTO_MAX_MB", "0.0001")  # ≈104 字节
        get_settings.cache_clear()

        resp = _传照片(client, FAKE_JPEG)  # 20 字节,稳在上限之内

        assert resp.status_code == 200, resp.json()

    def test_空body是400(self, client: TestClient) -> None:
        """前端拼请求时把文件漏了。**不许被当成"传了张空照片"收下** ——
        收下的话产物库里会多一份 0 字节的"照片",而复查那一栏点开是空的。"""
        之前 = _产物文件()

        resp = _传照片(client, b"")

        assert resp.status_code == 400
        assert resp.json()["error_code"] == "INVALID_INPUT"
        assert "照片" in resp.json()["user_msg"]
        assert _产物文件() == 之前

    def test_回执里没有documents键(self, client: TestClient) -> None:
        """它**不是动作端点、不出文书**,所以不在那份冻结的 ``data`` 形状里
        (模块头注「Envelope 的 data 形状是冻结的」那一节写明了这条例外)。

        钉死它是为了防"顺手补齐" —— 补一个恒空的 ``documents`` 上去,前端就会多渲
        一块永远空着的下载卡位;而 ``hazard_no`` / ``status`` 更不该有:
        这条端点压根不认识任何一条隐患,只是把字节存下来。
        """
        resp = _传照片(client, FAKE_JPEG)

        assert resp.status_code == 200
        assert set(resp.json()["data"]) == {"photo_id", "filename"}

    def test_魔数认出来的扩展名全在白名单里(self) -> None:
        """把那道导入期硬失败写成人看得见的形式(同 ``_deleted这个哨兵值`` 那条的手法)。

        漏配的表现:照片落盘没有后缀,而 ``_require_photo`` 只查 kind 与文件在不在、
        照样放行 —— 证据链里挂着一张**浏览器打不开**的照片,不会有任何报错。
        """
        assert set(supervision_api._SNIFFED_EXTS) <= ALLOWED_IMAGE_EXT

    def test_同一张照片传两次是两个编号(self, client: TestClient) -> None:
        """**刻意不做去重、不做幂等**(实现的 docstring 记了这笔账)。

        写成用例是为了让下一个人看见这是**决定**而不是遗漏:哪天真要上幂等,
        得先想清楚是加客户端 event_id(打卡那条链的账)还是按内容 hash 建索引,
        而不是看到"传两次两条"就当 bug 修。
        """
        第一次 = _传照片(client, FAKE_JPEG).json()["data"]["photo_id"]
        第二次 = _传照片(client, FAKE_JPEG).json()["data"]["photo_id"]

        assert 第一次 != 第二次
