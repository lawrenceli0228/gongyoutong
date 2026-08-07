"""Report Agent 的工具集 —— 把一次照片巡检渲染成可存档的巡检记录(docx)。

===========================================================================
数据保真:隐患内容**不经过任何 LLM 转抄**
---------------------------------------------------------------------------
巡检记录是要留档、可能用于追责的文档,里面的隐患项错一个字就是另一份文件。
所以本工具只收 artifact_id,内部**重新调用** safety 的 analyze_site_photo:

    render_inspection_report(artifact_id)
        │ analyze_site_photo(artifact_id)   ← 提示词已冻结 + 键算在图片内容上
        │                                     → 必然命中磁盘缓存:零成本,
        │                                       且与 safety 刚讲给用户的判断逐字节一致
        ▼ envelope.data(五键契约:label/violations/severity/max_severity/note)
    python-docx 渲染 → artifacts.register(REPORT) → 报告编号

英雄链里 safety 先跑(把缓存焐热),硬边到 report,report 的模型只负责把
32 位编号从上文搬进工具参数 —— 搬错一位立刻 NOT_FOUND 报错,
**不可能产出一份内容错误却看起来正常的文档**。这比让模型转抄隐患清单安全一个量级。

依赖方向说明:report → safety 的 import 是**有意的**,英雄链本身就是这条耦合
(拍照识违规 → 自动出记录)。别为了"解耦"把 analyze 抽到 core 去 ——
它的提示词、缓存维度、词表守卫全长在 safety 包里,搬家只会制造第二真相源。
"""

from __future__ import annotations

import io
import logging
from datetime import UTC, datetime
from typing import Any, Final

from docx import Document
from docx.shared import Pt
from langchain_core.tools import tool

from gyt.agents.safety.tools import analyze_site_photo
from gyt.config import get_settings
from gyt.core import artifacts
from gyt.core.artifacts import ArtifactKind
from gyt.core.errors import Envelope, ErrorCode, fail, ok, tool_guard

logger = logging.getLogger(__name__)

LABEL_ZH: Final[dict[str, str]] = {
    "violation": "发现安全隐患",
    "compliant": "未见明显隐患",
    "not_site": "非作业现场照片",
}
"""三态结论 → 文档用语。认不出的 label 原样落进文档并加「(待人工确认)」——
不吞、不猜,和 safety 工具的透传哲学一致。"""

_DISCLAIMER: Final[str] = (
    "本记录由工友通 AI 初筛生成,仅作快速筛查与留档之用;"
    "涉及精确测量与专业资质判断的事项,以持证安全员现场判定为准。"
)
"""落款免责句 —— 「初筛与记录」的定位直接写进每一份产物,这是技术说明边界①的落地。"""

_RENDER_DESCRIPTION = (
    "把一张已检查过的现场照片生成正式的巡检记录文档(Word)。"
    "参数 artifact_id 是照片的产物编号(32 位十六进制),从对话上文里原样复制。"
    "返回巡检记录的报告编号与隐患摘要。用户要「出巡检记录」「留档」「出报告」时调它。"
)


def _render_docx(
    *,
    photo_id: str,
    report_no: str,
    data: dict[str, Any],
    settings: Any,
) -> bytes:
    """按巡检记录模板渲染 docx,返回文件字节。纯函数式:不落盘、不改入参。"""
    label = str(data.get("label") or "")
    violations: list[str] = list(data.get("violations") or [])
    severity: dict[str, str] = dict(data.get("severity") or {})
    note = str(data.get("note") or "")

    doc = Document()
    doc.add_heading("工地安全巡检记录", level=0)

    meta = doc.add_table(rows=4, cols=2)
    meta.style = "Table Grid"
    rows = [
        ("记录编号", report_no),
        ("生成时间", datetime.now(UTC).astimezone().strftime("%Y-%m-%d %H:%M")),
        ("照片编号", photo_id),
        ("巡检结论", LABEL_ZH.get(label, f"{label or '(空)'}(待人工确认)")),
    ]
    for (key, value), row in zip(rows, meta.rows, strict=True):
        row.cells[0].text = key
        row.cells[1].text = value

    doc.add_heading("隐患明细", level=1)
    if violations:
        table = doc.add_table(rows=1 + len(violations), cols=3)
        table.style = "Table Grid"
        for cell, head in zip(table.rows[0].cells, ("序号", "隐患项", "级别"), strict=True):
            cell.text = head
        for index, item in enumerate(violations, start=1):
            cells = table.rows[index].cells
            cells[0].text = str(index)
            cells[1].text = item
            cells[2].text = severity.get(item, "待定级")
    else:
        doc.add_paragraph("本次巡检未发现受控清单内的隐患。")

    if note:
        doc.add_heading("现场备注与待复核事项", level=1)
        doc.add_paragraph(note)

    doc.add_heading("生成信息", level=1)
    doc.add_paragraph(
        f"识别模型:{settings.model_vision} / 提示词版本:{settings.prompt_version}。"
        "隐患判定与级别由系统按固定规则直出,未经人工改写。"
    )
    tail = doc.add_paragraph(_DISCLAIMER)
    tail.runs[0].font.size = Pt(9)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@tool("render_inspection_report", description=_RENDER_DESCRIPTION)
@tool_guard
async def render_inspection_report(artifact_id: str) -> Envelope:
    """生成一张照片的巡检记录 docx,登记为 REPORT 产物。

    artifact_id
      │ 编号不合法/照片不存在/识别失败 ──▶ 原样透传 safety 工具的失败信封
      ▼ analyze_site_photo(缓存必中,数据与 safety 口径逐字节一致)
    五键契约 data
      ▼ python-docx 渲染 + artifacts.register(REPORT)
    ok(data={report_id, filename, label, violations, severity, max_severity})
    """
    settings = get_settings()

    analysis = await analyze_site_photo.ainvoke({"artifact_id": (artifact_id or "").strip()})
    if not isinstance(analysis, dict) or not analysis.get("ok"):
        # 失败信封原样透传:里面的 user_msg 已经是中文人话(编号不对/照片没了/模型超时),
        # 在这里重新包装一层只会把「该怎么办」的信息越包越模糊。
        if isinstance(analysis, dict):
            return analysis
        return fail(ErrorCode.INTERNAL, detail=f"analyze 返回了非信封:{type(analysis).__name__}")

    data: dict[str, Any] = dict(analysis.get("data") or {})
    cleaned = (artifact_id or "").strip()
    # 记录编号用「GYT-日期-时刻」而不是产物编号:产物编号要注册后才有,而正文
    # 渲染在注册之前(注册需要文件字节)。时间戳号人念得出来、电话里报得清,
    # 产物编号(32 位 hex)在返回值里一并给出,两者都能唯一定位这份文件。
    report_no = datetime.now(UTC).astimezone().strftime("GYT-%Y%m%d-%H%M%S")

    payload = _render_docx(photo_id=cleaned, report_no=report_no, data=data, settings=settings)
    filename = f"巡检记录_{report_no}.docx"
    report_id = artifacts.register(payload, kind=ArtifactKind.REPORT, original_name=filename)
    stored = artifacts.resolve(report_id)
    logger.info("巡检记录已生成:%s(%s,%d 字节)", report_no, report_id, len(payload))

    violations = list(data.get("violations") or [])
    label = str(data.get("label") or "")
    if violations:
        summary = (
            f"巡检记录已生成(编号 {report_no}),"
            f"共 {len(violations)} 处隐患,最高级别 {data.get('max_severity')}。"
        )
    elif label == "not_site":
        summary = f"巡检记录已生成(编号 {report_no}),但照片被判定为非作业现场,请核对后重拍。"
    else:
        summary = f"巡检记录已生成(编号 {report_no}),本次未发现受控清单内的隐患。"

    return ok(
        data={
            "report_id": report_id,
            "report_no": report_no,
            "filename": filename,
            "path": str(stored),
            "label": label,
            "violations": violations,
            "severity": dict(data.get("severity") or {}),
            "max_severity": data.get("max_severity"),
        },
        user_msg=summary,
    )


REPORT_TOOLS: list = [render_inspection_report]

__all__ = ["LABEL_ZH", "REPORT_TOOLS", "render_inspection_report"]
