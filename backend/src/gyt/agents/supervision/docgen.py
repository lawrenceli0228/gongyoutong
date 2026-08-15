"""共用 docx 骨架 —— 六种监理文书 + 巡检记录,全仓只有这一处画 Word 文档。

===========================================================================
骨架长什么样
---------------------------------------------------------------------------
    标题(level 0)
    元信息表(键值两列,Table Grid)
    节 × N   ── 每节 = 小标题(level 1) + 一段正文  或  小标题 + 一张表
    签字栏(留空,D15)
    落款免责句(小一号字)

「填什么」由调用方给,「长什么样」由这里定。这么切的理由是 D8:六种文书的
差别全在内容,版式一模一样;各写各的渲染函数意味着将来加一行页眉要改六遍,
而漏改的那一份不会报错 —— 它只是长得跟别的不一样,而且没人会去比对。

===========================================================================
纯函数式:**返回 bytes,不落盘,不改入参**
---------------------------------------------------------------------------
这不是洁癖,是方案 §6.4「原子性 —— 文件先落盘,库后写」的前提:

    ① 渲染三份 docx 到内存(纯函数,不落盘)
    ② artifacts.register ×3 → 拿到三个 artifact_id   ← 文件写失败:整个操作
                                                        fail,库一行没动,零孤儿行
    ③ 一个事务:写 hazard_docs ×3 + 改 hazards.status ← 库写失败:留三个孤儿
                                                        文件,可接受

渲染阶段一旦自己落盘,①②之间就有了半成品:第二份渲染炸掉时第一份已经躺在
磁盘上,而它既没进注册表也没进库,连清理器都不认识它。所以本模块**不碰
文件系统** —— 落盘统一归 ``core/artifacts.py``。

同样地,入参一律只读:文书内容是要留档、可能用于追责的材料,渲染顺手改掉
调用方手里那份数据的话,后面写库写进去的就是被改过的版本,而两边对不上要
到有人打开文件核对时才发现。

===========================================================================
为什么住在 supervision 包里,而 report 也来 import 它
---------------------------------------------------------------------------
方案 §11 的 S2 泳道就是这么划的(``core/``、``supervision/docgen.py``、
``report/tools.py`` 一起改)。当前只有两个消费方,放在主用方包里比先搬进
core 更实在。**要是出现第三个消费方,就该搬到 ``core/`` 去** —— 那时
「agents/report import agents/supervision」这条横向边才真的开始碍事。
"""

from __future__ import annotations

import io
from collections.abc import Sequence
from dataclasses import dataclass
from typing import Final

from docx import Document
from docx.document import Document as DocxDocument
from docx.shared import Pt

# ``docx.Document`` 是**工厂函数**不是类,拿它当类型注解是错的(有
# ``from __future__ import annotations`` 兜着不会当场炸,但类型检查器读不懂)。
# 真正的类在 ``docx.document.Document``,所以这里分成两个名字导入。

# --- 版式常量(禁止在函数体里散落魔法值)-----------------------------------

_TITLE_LEVEL: Final[int] = 0
"""标题用 level 0(python-docx 的 Title 样式),与巡检记录原有版式一致。"""

_SECTION_LEVEL: Final[int] = 1
"""节的小标题用 level 1。"""

_TABLE_STYLE: Final[str] = "Table Grid"
"""带框线的表格样式。留档文件要能打印出来判读,无框线的表复印一次就散了。"""

_TAIL_FONT_PT: Final[int] = 9
"""免责句字号,比正文小一号 —— 沿用巡检记录原有版式。"""

SUPERVISION_DISCLAIMER: Final[str] = (
    "本文书由工友通 AI 依据现场影像与监理规则自动出稿,"
    "须经总监理工程师签字盖章后方为正式文件;"
    "签字之前不得据以停工、复工或对外发出。"
)
"""**监理文书**的落款免责句(D15:「AI 出稿,经总监理工程师签字后生效」)。

D15 把这六种文书的定位从「演示件」抬到了「正式文书的初稿」,那就必须在每一份
产物上把边界写死:未签字的这份**不是**正式文件。写「不得据以停工、复工或
对外发出」而不是笼统的「仅供参考」,是因为这三件事正是这几种文书会引发的
实际动作 —— 说清楚哪几件事不能干,比说「仅供参考」有用。

⚠️ **它和巡检记录那句不是一回事,不许互相顶替。**巡检记录的免责句
(``agents/report/tools.py`` 的 ``_DISCLAIMER``)讲的是「AI 初筛 vs 持证安全员
现场判定」,压根没有「总监签字后生效」这一环。结构上的保证是
``render_document`` 的 ``disclaimer`` **必填且没有默认值** —— 谁都不可能
「漏传一个参数,于是悄悄套上了另一句」。
"""

SUPERVISION_SIGNATURE_LINES: Final[tuple[str, ...]] = (
    "总监理工程师(签字):",
    "项目监理机构(盖章):",
    "日期:    年    月    日",
)
"""监理文书的签字栏。**冒号后面一律留空**(D15)。

不许自动填上任何名字或日期:签字这个动作本身就是「有个人认下了这份文书」的
唯一凭据,系统替他填等于伪造。留空的横线由签字的人自己写。
"""


@dataclass(frozen=True)
class TextSection:
    """一个节:小标题 + 一段正文。

    ``body`` 为空串 = 只出小标题不出正文(用于「本节无内容但版式要在」的场合)。
    """

    heading: str
    body: str = ""


@dataclass(frozen=True)
class TableSection:
    """一个节:小标题 + 一张表。

    ``rows`` 为空时不画空表,改画 ``empty_note`` 那句话 —— 一张只有表头的空表
    在留档文件里读起来像「这一栏还没填」,而「本次未发现」是个明确结论,两者
    在追责场合意思完全不同。``empty_note`` 也为空则这一节只剩小标题。

    ``header`` / ``rows`` 用**元组**不用列表:frozen 只冻结字段绑定,字段里
    塞个列表照样能被外面就地改掉,而这是留档材料。
    """

    heading: str
    header: tuple[str, ...]
    rows: tuple[tuple[str, ...], ...] = ()
    empty_note: str = ""


Section = TextSection | TableSection
"""一个节的两种形态。加第三种形态时记得在 ``render_document`` 的分支里加档 ——
那里的 ``else`` 是显式 ``TypeError``,不会静默跳过。"""


def _render_table(doc: DocxDocument, section: TableSection) -> None:
    """把一个表格节画进文档。校验先做完再建表,免得留下画到一半的残表。"""
    if not section.header:
        raise ValueError(f"「{section.heading}」是表格节却没有表头 —— 没有表头的表没法判读")

    width = len(section.header)
    for index, values in enumerate(section.rows, start=1):
        if len(values) != width:
            raise ValueError(
                f"「{section.heading}」第 {index} 行有 {len(values)} 格,表头是 {width} 列 —— "
                "列数对不上会把内容错位到相邻列里,而错位的留档文书从外观上看不出问题"
            )

    table = doc.add_table(rows=1 + len(section.rows), cols=width)
    table.style = _TABLE_STYLE
    for cell, head in zip(table.rows[0].cells, section.header, strict=True):
        cell.text = head
    for row, values in zip(table.rows[1:], section.rows, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value


def render_document(
    *,
    title: str,
    meta: Sequence[tuple[str, str]] = (),
    sections: Sequence[Section] = (),
    disclaimer: str,
    signature_lines: Sequence[str] = (),
) -> bytes:
    """按共用骨架渲染一份 docx,返回文件字节。纯函数式:不落盘、不改入参。

    参数:
        title:           文档标题。
        meta:            元信息表,``(键, 值)`` 有序对;空则不画表。
        sections:        正文各节,按给的顺序出。
        disclaimer:      落款免责句。**必填、无默认值**,理由见
                         ``SUPERVISION_DISCLAIMER`` 的说明(防两句话互相顶替)。
        signature_lines: 签字栏各行;监理文书传 ``SUPERVISION_SIGNATURE_LINES``,
                         巡检记录不传(它没有签字这一环)。

    标题与免责句空着就抛 ``ValueError``:一份没标题的文书没法归档,一份没有
    免责句的文书等于对外宣称「这是终稿」—— 两件都不是能靠默认值糊过去的事。
    """
    if not title.strip():
        raise ValueError("文书标题不能为空 —— 归档时全靠它认这是什么文件")
    if not disclaimer.strip():
        raise ValueError("落款免责句不能为空 —— 它是「AI 出稿、签字后生效」这条边界的唯一落地点")

    doc = Document()
    doc.add_heading(title, level=_TITLE_LEVEL)

    if meta:
        table = doc.add_table(rows=len(meta), cols=2)
        table.style = _TABLE_STYLE
        for (key, value), row in zip(meta, table.rows, strict=True):
            row.cells[0].text = key
            row.cells[1].text = value

    for section in sections:
        # **先认类型,再动笔。** 认不出的东西要在写下小标题之前就炸掉:
        # 反过来的话文档里会留一个没有内容的孤零零小标题,而异常一旦被上层
        # 吞掉(比如 tool_guard),那份缺了一整节的文书就这么出去了 ——
        # 悄悄少一节的文书,外观上完全正常。
        if not isinstance(section, TextSection | TableSection):
            raise TypeError(
                f"认不出的节类型 {type(section).__name__},只收 TextSection / TableSection"
            )

        doc.add_heading(section.heading, level=_SECTION_LEVEL)
        if isinstance(section, TextSection):
            if section.body:
                doc.add_paragraph(section.body)
        elif section.rows:
            _render_table(doc, section)
        elif section.empty_note:
            doc.add_paragraph(section.empty_note)

    for line in signature_lines:
        doc.add_paragraph(line)

    tail = doc.add_paragraph(disclaimer)
    tail.runs[0].font.size = Pt(_TAIL_FONT_PT)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


__all__ = [
    "SUPERVISION_DISCLAIMER",
    "SUPERVISION_SIGNATURE_LINES",
    "Section",
    "TableSection",
    "TextSection",
    "render_document",
]
